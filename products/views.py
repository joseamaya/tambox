# -*- coding: utf-8 -*- 
import logging

from django.views.generic.base import View, TemplateView
from django.views.generic.list import ListView
from django.views.generic.edit import FormView, UpdateView, CreateView
from django.urls import reverse_lazy, reverse
from django.http.response import HttpResponseRedirect
import json
from django.http import HttpResponse
import simplejson
from openpyxl import Workbook
from django.views.generic.detail import DetailView
from security.permissions import requires
from django.utils.decorators import method_decorator
from accounting.forms import UploadForm
from django.shortcuts import render
from products.models import Product, UnitOfMeasure, ProductGroup
from products.forms import ProductGroupForm, ProductForm, ServiceForm,\
    UnitOfMeasureForm
from accounting.models import Account, StockType
from tambox.views import CsvImportMixin, AjaxOnlyMixin

logger = logging.getLogger(__name__)


class Dashboard(View):

    def get(self, request, *args, **kwargs):
        notification_list = []
        product_count = Product.objects.filter(is_service=False).count()
        unit_of_measure_count = UnitOfMeasure.objects.count()
        supply_group_count = ProductGroup.objects.count()
        service_count = Product.objects.filter(is_service=True).count()
        unit_of_measure, creado = UnitOfMeasure.objects.get_or_create(code='SERV',
                                                                   defaults={'description': 'SERVICIO'})
        if creado:
            notification_list.append("Se ha creado la unidad de medida SERVICIO")
        if product_count == 0:
            notification_list.append("No se ha creado ningún producto")
        if unit_of_measure_count == 0:
            notification_list.append("No se ha creado ningún tipo de unidad de medida")
        if supply_group_count == 0:
            notification_list.append("No se ha creado ningún grupo de productos")
        if service_count == 0:
            notification_list.append("No se ha creado ningún service")
        context = {'notifications': notification_list}
        return render(request, 'products/products_dashboard.html', context)


class ProductDescriptionSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('description', 'search_type')

    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            description = request.GET['description']
            search_type = request.GET['search_type']
            if search_type == 'TODOS':
                products = Product.objects.filter(description__icontains=description).select_related(
                    'unit_of_measure').order_by('description')[:20]
            elif search_type == 'PRODUCTOS':
                products = Product.objects.filter(description__icontains=description,
                                                    is_service=False).select_related(
                    'unit_of_measure').order_by('description')[:20]
            elif search_type == 'SERVICIOS':
                products = Product.objects.filter(description__icontains=description,
                                                    is_service=True).select_related(
                    'unit_of_measure').order_by('description')[:20]

            product_list = []
            for product in products:
                product_json = {}
                product_json['label'] = product.description
                product_json['code'] = product.code
                product_json['description'] = product.description
                product_json['unit'] = product.unit_of_measure.description
                product_json['price'] = str(product.price)
                product_list.append(product_json)
            data = json.dumps(product_list)
            return HttpResponse(data, 'application/json')


class ProductCodeSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('code',)

    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.GET['code']
            products = Product.objects.filter(code__icontains=code).select_related('unit_of_measure')[:20]
            product_list = []
            for product in products:
                product_json = {}
                product_json['label'] = product.code
                product_json['code'] = product.code
                product_json['description'] = product.description
                product_json['unit'] = product.unit_of_measure.description
                product_json['price'] = str(product.price)
                product_list.append(product_json)
            data = json.dumps(product_list)
            return HttpResponse(data, 'application/json')


class ProductGroupImport(CsvImportMixin, FormView):
    template_name = 'products/product_group_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('products:product_group_list')

    def process_row(self, row):
        try:
            account_number = Account.objects.get(account_number=row[0])
            ProductGroup.objects.get_or_create(description=row[1],
                                                 defaults={'account': account_number})
        except Account.DoesNotExist:
            pass


class ServiceImport(CsvImportMixin, FormView):
    template_name = 'products/service_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('products:service_list')

    def form_valid(self, form):
        try:
            return super(ServiceImport, self).form_valid(form)
        except ProductGroup.DoesNotExist:
            return HttpResponseRedirect(reverse('products:product_group_create'))

    def process_row(self, row):
        group = ProductGroup.objects.get(code=row[0].strip())
        Product.objects.get_or_create(description=row[1],
                                       defaults={'product_group': group,
                                                 'is_service': True})


class ProductImport(CsvImportMixin, FormView):
    template_name = 'products/product_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('products:product_list')

    def process_row(self, row):
        try:
            group = ProductGroup.objects.get(code=row[0].strip())
            unit_code = row[2][0:5]
            und, creado = UnitOfMeasure.objects.get_or_create(code=unit_code.strip(),
                                                             defaults={'code': unit_code,
                                                                       'description': row[2].strip()})
            if row[3] != '':
                price = row[3]
            else:
                price = 0
            stock_type = StockType.objects.get(sunat_code=row[4].strip())
            product, creado = Product.objects.get_or_create(description=row[1].strip(),
                                                              defaults={'unit_of_measure': und,
                                                                        'product_group': group,
                                                                        'price': price,
                                                                        'stock_type': stock_type})
        except Exception:
            logger.warning("No se pudo importar el producto %s", row[1], exc_info=True)


class ProductStockQuery(AjaxOnlyMixin, TemplateView):

    required_params = ('code',)

    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.GET['code']
            product = Product.objects.get(code=code)
            product_json = {}
            product_json['stock'] = product.stock
            data = simplejson.dumps(product_json)
            return HttpResponse(data, 'application/json')


class ProductGroupCreate(CreateView):
    model = ProductGroup
    template_name = 'products/product_group_form.html'
    form_class = ProductGroupForm
    success_url = reverse_lazy('products:product_group_list')

    @method_decorator(requires('products.add_productgroup'))
    def dispatch(self, *args, **kwargs):
        return super(ProductGroupCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:product_group_detail', args=[self.object.pk])


class ProductCreate(CreateView):
    model = Product
    context_object_name = 'product'
    template_name = 'products/product_form.html'
    form_class = ProductForm

    @method_decorator(requires('products.add_product'))
    def dispatch(self, *args, **kwargs):
        return super(ProductCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:product_detail', args=[self.object.pk])


class UnitOfMeasureCreate(CreateView):
    template_name = 'products/unit_of_measure_form.html'
    form_class = UnitOfMeasureForm

    @method_decorator(requires('products.add_unitofmeasure'))
    def dispatch(self, *args, **kwargs):
        return super(UnitOfMeasureCreate, self).dispatch(*args, **kwargs)

    def form_valid(self, form):
        self.object = form.save(commit=False)
        self.object.changed_by = self.request.user
        self.object.save()
        return super(UnitOfMeasureCreate, self).form_valid(form)

    def get_success_url(self):
        return reverse('products:unit_of_measure_detail', args=[self.object.pk])


class ServiceCreate(CreateView):
    template_name = 'products/service_form.html'
    form_class = ServiceForm

    @method_decorator(requires('products.add_product'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:service_detail', args=[self.object.code])


class ProductDetail(DetailView):
    model = Product
    template_name = 'products/product_detail.html'


class ProductGroupDetail(DetailView):
    model = ProductGroup
    template_name = 'products/product_group_detail.html'


class UnitOfMeasureDetail(DetailView):
    model = UnitOfMeasure
    template_name = 'products/unit_of_measure_detail.html'


class ServiceDetail(DetailView):
    model = Product
    template_name = 'products/service_detail.html'


class UnitOfMeasureDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('products.delete_unitofmeasure'))
    def dispatch(self, *args, **kwargs):
        return super(UnitOfMeasureDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            id = request.POST['id']
            unit_of_measure = UnitOfMeasure.objects.get(pk=id)
            unit_of_measure_json = {}
            unit_of_measure_json['unit'] = unit_of_measure.code
            if len(unit_of_measure.products.all()) > 0:
                unit_of_measure_json['productos'] = 'SI'
            else:
                unit_of_measure_json['productos'] = 'NO'
                UnitOfMeasure.objects.filter(pk=id).update(is_active=False)
            data = simplejson.dumps(unit_of_measure_json)
            return HttpResponse(data, 'application/json')


class ProductGroupDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('products.delete_productgroup'))
    def dispatch(self, *args, **kwargs):
        return super(ProductGroupDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            product_group = ProductGroup.objects.get(pk=code)
            product_group_json = {}
            product_group_json['code'] = product_group.code
            product_group_json['description'] = product_group.description
            if len(product_group.products.all()) > 0:
                product_group_json['productos'] = 'SI'
            else:
                product_group_json['productos'] = 'NO'
                ProductGroup.objects.filter(pk=code).update(is_active=False)
            data = simplejson.dumps(product_group_json)
            return HttpResponse(data, 'application/json')


class ProductDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('products.delete_product'))
    def dispatch(self, *args, **kwargs):
        return super(ProductDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            product = Product.objects.get(pk=code)
            product_json = {}
            product_json['code'] = product.code
            product_json['description'] = product.description
            if (product.order_details.exists() or product.purchase_order_details.exists()
                    or product.service_order_details.exists()
                    or product.movement_details.exists()
                    or product.requirement_details.exists()):
                product_json['relaciones'] = 'SI'
            else:
                product_json['relaciones'] = 'NO'
                Product.objects.filter(pk=code).update(is_active=False)
            data = simplejson.dumps(product_json)
            return HttpResponse(data, 'application/json')


class ServiceDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('products.delete_product'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            service = Product.objects.get(code=code)
            service_json = {}
            service_json['code'] = code
            if len(service.service_order_details.all()) > 0:
                service_json['orders'] = 'SI'
            else:
                service_json['orders'] = 'NO'
                Product.objects.filter(code=code).update(is_active=False)
            data = simplejson.dumps(service_json)
            return HttpResponse(data, 'application/json')


class UnitOfMeasureList(ListView):
    model = UnitOfMeasure
    template_name = 'products/unit_of_measure_list.html'
    context_object_name = 'units'
    queryset = UnitOfMeasure.objects.filter(is_active=True).order_by('description')

    @method_decorator(
        requires('products.ver_tabla_unidades_medida'))
    def dispatch(self, *args, **kwargs):
        return super(UnitOfMeasureList, self).dispatch(*args, **kwargs)


class ServiceList(ListView):
    model = Product
    template_name = 'products/service_list.html'
    context_object_name = 'services'
    queryset = Product.objects.filter(is_active=True, is_service=True).order_by('description')

    @method_decorator(requires('products.ver_tabla_productos'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceList, self).dispatch(*args, **kwargs)


class ProductGroupList(ListView):
    model = ProductGroup
    template_name = 'products/product_group_list.html'
    context_object_name = 'product_groups'
    queryset = ProductGroup.objects.filter(is_active=True).order_by('code')

    @method_decorator(
        requires('products.ver_tabla_grupos_productos'))
    def dispatch(self, *args, **kwargs):
        return super(ProductGroupList, self).dispatch(*args, **kwargs)


class ProductList(ListView):
    model = Product
    template_name = 'products/product_list.html'
    context_object_name = 'products'
    queryset = Product.objects.filter(is_service=False, is_active=True).order_by('code')

    @method_decorator(requires('products.ver_tabla_productos'))
    def dispatch(self, *args, **kwargs):
        return super(ProductList, self).dispatch(*args, **kwargs)


class ProductListByGroup(ListView):
    model = Product
    template_name = 'products/product_list.html'
    context_object_name = 'products'

    @method_decorator(requires('products.ver_tabla_productos'))
    def dispatch(self, *args, **kwargs):
        return super(ProductListByGroup, self).dispatch(*args, **kwargs)

    def get_queryset(self):
        group = ProductGroup.objects.get(pk=self.kwargs['group'])
        queryset = group.products.all()
        return queryset


class ProductUpdate(UpdateView):
    model = Product
    context_object_name = 'product'
    template_name = 'products/product_form.html'
    form_class = ProductForm

    @method_decorator(requires('products.change_product'))
    def dispatch(self, *args, **kwargs):
        return super(ProductUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:product_detail', args=[self.object.pk])


class UnitOfMeasureUpdate(UpdateView):
    model = UnitOfMeasure
    template_name = 'products/unit_of_measure_form.html'
    form_class = UnitOfMeasureForm

    @method_decorator(requires('products.change_unitofmeasure'))
    def dispatch(self, *args, **kwargs):
        return super(UnitOfMeasureUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:unit_of_measure_detail', args=[self.object.pk])


class ProductGroupUpdate(UpdateView):
    model = ProductGroup
    template_name = 'products/product_group_form.html'
    form_class = ProductGroupForm
    success_url = reverse_lazy('products:product_group_list')

    @method_decorator(
        requires('products.change_productgroup'))
    def dispatch(self, *args, **kwargs):
        return super(ProductGroupUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:product_group_detail', args=[self.object.pk])


class ServiceUpdate(UpdateView):
    model = Product
    template_name = 'products/service_form.html'
    form_class = ServiceForm

    @method_decorator(requires('products.change_product'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('products:service_detail', args=[self.object.pk])


class ProductExcelReport(TemplateView):

    def get(self, request, *args, **kwargs):
        products = Product.objects.filter(is_active=True).order_by('code')
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE PRODUCTOS'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'CODIGO'
        ws['C3'] = 'DESCRIPCION'
        ws['D3'] = 'DESCR_ABREV'
        ws['E3'] = 'GRUPO'
        ws['F3'] = 'UNIDAD'
        ws['G3'] = 'MARCA'
        ws['H3'] = 'MODELO'
        ws['I3'] = 'PRECIO'
        ws['J3'] = 'CREADO'
        cont = 4
        for product in products:
            ws.cell(row=cont, column=2).value = product.code
            ws.cell(row=cont, column=3).value = product.description
            ws.cell(row=cont, column=4).value = product.desc_abreviada
            ws.cell(row=cont, column=5).value = product.product_group.description
            ws.cell(row=cont, column=6).value = product.unit_of_measure.description
            ws.cell(row=cont, column=7).value = product.brand
            ws.cell(row=cont, column=8).value = product.model
            ws.cell(row=cont, column=9).value = product.price
            ws.cell(row=cont, column=9).number_format = '#.00000'
            ws.cell(row=cont, column=10).value = product.created
            ws.cell(row=cont, column=10).number_format = 'dd/mm/yyyy hh:mm:ss'
            cont = cont + 1
        file_name = "ProductList.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class ProductGroupExcelReport(TemplateView):

    def get(self, request, *args, **kwargs):
        product_groups = ProductGroup.objects.filter(is_active=True).order_by('code')
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE GRUPOS DE PRODUCTOS'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'CODIGO'
        ws['C3'] = 'DESCRIPCION'
        ws['D3'] = 'CTA_CONTABLE'
        ws['E3'] = 'CREADO'
        cont = 4
        for product_group in product_groups:
            ws.cell(row=cont, column=2).value = product_group.code
            ws.cell(row=cont, column=3).value = product_group.description
            ws.cell(row=cont, column=4).value = product_group.account.account_number
            ws.cell(row=cont, column=5).value = product_group.created
            ws.cell(row=cont, column=5).number_format = 'dd/mm/yyyy hh:mm:ss'
            cont = cont + 1
        file_name = "ProductGroupList.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class UnitOfMeasureExcelReport(TemplateView):

    def get(self, request, *args, **kwargs):
        units = UnitOfMeasure.objects.filter(is_active=True).order_by('code')
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE ONES DE MEDIDA'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'UNIDAD'
        ws['C3'] = 'DESCRIPCIÓN'
        ws['D3'] = 'ESTADO'
        cont = 4
        for unit in units:
            ws.cell(row=cont, column=2).value = unit.code
            ws.cell(row=cont, column=3).value = unit.description
            ws.cell(row=cont, column=4).value = unit.is_active
            cont = cont + 1
        file_name = "UnidadesMedida.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class ServiceExcelReport(TemplateView):

    def get(self, request, *args, **kwargs):
        services = Product.objects.filter(is_service=True, is_active=True).order_by('code')
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE SERVICIOS'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'CODIGO'
        ws['C3'] = 'DESCRIPCION'
        ws['D3'] = 'ESTADO'
        cont = 4
        for service in services:
            ws.cell(row=cont, column=2).value = service.code
            ws.cell(row=cont, column=3).value = service.description
            ws.cell(row=cont, column=4).value = service.is_active
            cont = cont + 1
        file_name = "ServiceList.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response
