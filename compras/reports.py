# -*- coding: utf-8 -*-

from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
from reportlab.platypus import Table
from reportlab.lib import colors
# from reportlab.lib.pagesizes import cm
from reportlab.platypus.flowables import Spacer, ListFlowable
from reportlab.graphics.shapes import Line, Drawing
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from django.conf import settings
import os
from io import BytesIO
from administracion.models import Position
from compras.models import PurchaseOrderDetail, ServiceOrderDetail, ServiceConformityDetail
from django.core.exceptions import ObjectDoesNotExist
from django.utils.encoding import force_str
from tambox.config import company, configuration
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import Side


class PurchaseOrderReport():

    def __init__(self, pagesize, orden_compra):
        self.orden_compra = orden_compra
        self.buffer = BytesIO()
        if pagesize == 'A4':
            self.pagesize = A4
        elif pagesize == 'Letter':
            self.pagesize = letter
        self.width, self.height = self.pagesize

    def header_table(self, styles):
        orden_compra = self.orden_compra
        sp = ParagraphStyle('parrafos',
                            alignment=TA_CENTER,
                            fontSize=14,
                            fontName="Times-Roman")
        try:
            archivo_imagen = os.path.join(settings.MEDIA_ROOT, str(company().logo))
            image = Image(archivo_imagen, width=90, height=50, hAlign='LEFT')
        except Exception:
            image = Paragraph(u"LOGO", sp)

        nro = Paragraph(u"ORDEN DE COMPRA", sp)
        tax_id = Paragraph("R.U.C." + company().tax_id, sp)
        encabezado = [[image, nro, tax_id], ['', u"N°" + orden_compra.code,
                                           company().district + " " + orden_compra.date.strftime('%d de %b de %Y')]]
        header_table = Table(encabezado, colWidths=[4 * cm, 9 * cm, 6 * cm])
        header_table.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (2, 1), 'CENTER'),
                ('VALIGN', (0, 0), (2, 0), 'CENTER'),
                ('VALIGN', (1, 1), (2, 1), 'TOP'),
                ('SPAN', (0, 0), (0, 1)),

            ]
        ))
        return header_table

    def data_table(self, styles):
        order = self.orden_compra
        izquierda = ParagraphStyle('parrafos',
                                   alignment=TA_LEFT,
                                   fontSize=10,
                                   fontName="Times-Roman")
        quotation = order.quotation
        if quotation is None:
            supplier = order.supplier
        else:
            supplier = order.quotation.supplier
        razon_social_proveedor = Paragraph(u"SEÑOR(ES): " + supplier.business_name, izquierda)
        ruc_proveedor = Paragraph(u"R.U.C.: " + supplier.tax_id, izquierda)
        address = Paragraph(u"DIRECCIÓN: " + supplier.address, izquierda)
        try:
            phone = Paragraph(u"TELÉFONO: " + supplier.phone, izquierda)
        except TypeError:
            phone = Paragraph(u"TELÉFONO: -", izquierda)
        try:
            reference = Paragraph(
                u"REFERENCIA: " + order.quotation.requirement.code + " - " + order.quotation.requirement.office.name,
                izquierda)
        except (ObjectDoesNotExist, AttributeError):
            reference = Paragraph(u"REFERENCIA: ", izquierda)
        process = Paragraph(u"PROCESO: " + order.process, izquierda)
        nota = Paragraph(u"Sírvase remitirnos según especificaciones que detallamos lo next: ", izquierda)
        datos = [[razon_social_proveedor, ruc_proveedor], [address, phone], [reference, ''], [process, ''],
                 [nota, '']]
        tabla_detalle = Table(datos, colWidths=[11 * cm, 9 * cm])
        tabla_detalle.setStyle(TableStyle(
            [
                ('SPAN', (0, 2), (1, 2)),
            ]
        ))
        return tabla_detalle

    def tabla_detalle(self):
        order = self.orden_compra
        encabezados = ['Item', 'Cantidad', 'Unidad', u'Descripción', 'Precio', 'Total']
        detalles = PurchaseOrderDetail.objects.filter(order=order).order_by('pk')
        sp = ParagraphStyle('parrafos')
        sp.alignment = TA_JUSTIFY
        sp.fontSize = 8
        sp.fontName = "Times-Roman"
        lista_detalles = []
        for detail in detalles:
            try:
                tupla_producto = [Paragraph(str(detail.line_number), sp),
                                  Paragraph(str(detail.quantity), sp),
                                  Paragraph(
                                      detail.quotation_detail.requirement_detail.product.unit_of_measure.description,
                                      sp),
                                  Paragraph(detail.quotation_detail.requirement_detail.product.description, sp),
                                  Paragraph(str(detail.price), sp),
                                  Paragraph(str(detail.amount), sp)]
            except (ObjectDoesNotExist, AttributeError):
                tupla_producto = [Paragraph(str(detail.line_number), sp),
                                  Paragraph(str(detail.quantity), sp),
                                  Paragraph(detail.product.unit_of_measure.description, sp),
                                  Paragraph(detail.product.description, sp),
                                  Paragraph(str(detail.price), sp),
                                  Paragraph(str(detail.amount), sp)]
            lista_detalles.append(tupla_producto)
        adicionales = [('', '', '', '', '')] * (15 - len(lista_detalles))
        tabla_detalle = Table([encabezados] + lista_detalles + adicionales,
                              colWidths=[0.8 * cm, 2 * cm, 2.5 * cm, 10.2 * cm, 2 * cm, 2.5 * cm])
        style = TableStyle(
            [
                ('ALIGN', (0, 0), (4, 0), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('ALIGN', (4, 1), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]
        )
        tabla_detalle.setStyle(style)
        return tabla_detalle

    def total_in_words_table(self):
        order = self.orden_compra
        total_in_words = [("SON: " + order.total_in_words, '')]
        total_in_words_table = Table(total_in_words, colWidths=[17.5 * cm, 2.5 * cm])
        total_in_words_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (1, 0), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        return total_in_words_table

    def others_table(self):
        order = self.orden_compra
        p = ParagraphStyle('parrafos',
                           alignment=TA_CENTER,
                           fontSize=8,
                           fontName="Times-Roman")
        sub_total = Paragraph(u"SUBTOTAL: ", p)
        igv = Paragraph(u"IGV: ", p)
        total = Paragraph(u"TOTAL: ", p)
        datos_otros = [
            [Paragraph(u"LUGAR DE ENTREGA", p), Paragraph(u"PLAZO DE ENTREGA", p), Paragraph(u"FORMA DE PAGO", p),
             sub_total, order.subtotal],
            [Paragraph(company().address(), p), Paragraph(u"INMEDIATA", p), Paragraph(order.payment_method.description, p),
             igv, str(order.igv)],
            ['', '', '', total, str(order.total)],
            ]
        others_table = Table(datos_otros, colWidths=[5.5 * cm, 5 * cm, 5 * cm, 2 * cm, 2.5 * cm])
        others_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (2, 2), 1, colors.black),
                ('SPAN', (0, 1), (0, 2)),
                ('SPAN', (1, 1), (1, 2)),
                ('SPAN', (2, 1), (2, 2)),
                ('GRID', (4, 0), (4, 2), 1, colors.black),
                ('VALIGN', (0, 1), (2, 1), 'MIDDLE'),
            ]
        ))
        return others_table

    def notes_table(self):
        order = self.orden_compra
        p = ParagraphStyle('parrafos',
                           alignment=TA_JUSTIFY,
                           fontSize=8,
                           fontName="Times-Roman")
        obs = Paragraph("OBSERVACIONES: " + order.notes, p)
        notes = [[obs]]
        notes_table = Table(notes, colWidths=[20 * cm], rowHeights=1.8 * cm)
        notes_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (0, 2), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]
        ))
        return notes_table

    def budget_impact_table(self):
        p = ParagraphStyle('parrafos',
                           alignment=TA_JUSTIFY,
                           fontSize=8,
                           fontName="Times-Roman")
        hoja_afectacion = Paragraph(u"HOJA DE AFECTACIÓN PRESUPUESTAL: ", p)
        importante = Paragraph(u"IMPORTANTE: ", p)
        recibido = Paragraph(u"RECIBIDO POR: ", p)
        signature = Paragraph(u"FIRMA: ", p)
        name = Paragraph(u"NOMBRE: ", p)
        dni = Paragraph(u"DNI: ", p)
        lista = ListFlowable([
            Paragraph("""Consignar el número de la presente Orden de Compra en su Guía de Remisión y Factura. 
                          Facturar a nombre de """ + force_str(company().business_name), p),
            Paragraph("El " + force_str(company().business_name) + """, se reserva el derecho de devolver 
                          la mercaderia, sino se ajusta a las especificaciones requeridas, asimismo de anular la presente 
                          Orden de Compra.""", p),
            Paragraph("""El pago de toda factura se hará de acuerdo a las condiciones establecidas.""", p)
        ], bulletType='1'
        )
        datos_otros = [[hoja_afectacion, ''],
                       [importante, recibido],
                       [lista, ''],
                       ['', signature],
                       ['', name],
                       ['', dni],
                       ]
        budget_impact_table = Table(datos_otros, colWidths=[10 * cm, 10 * cm])
        budget_impact_table.setStyle(TableStyle(
            [
                ('ALIGN', (0, 1), (1, 1), 'CENTER'),
                ('SPAN', (0, 2), (0, 5)),
            ]
        ))
        return budget_impact_table

    def render(self):
        y = 300
        buffer = self.buffer
        doc = SimpleDocTemplate(buffer,
                                rightMargin=50,
                                leftMargin=50,
                                topMargin=20,
                                bottomMargin=50,
                                pagesize=self.pagesize)

        elements = []
        styles = getSampleStyleSheet()
        elements.append(self.header_table(styles))
        elements.append(Spacer(1, 0.25 * cm))
        elements.append(self.data_table(styles))
        elements.append(Spacer(1, 0.25 * cm))
        elements.append(self.tabla_detalle())
        elements.append(Spacer(1, 0.25 * cm))
        elements.append(self.total_in_words_table())
        elements.append(Spacer(1, 0.25 * cm))
        elements.append(self.others_table())
        elements.append(Spacer(1, 0.25 * cm))
        elements.append(self.notes_table())
        elements.append(Spacer(1, 0.25 * cm))
        elements.append(self.budget_impact_table())
        linea_firma = Line(280, y - 250, 470, y - 250)
        d = Drawing(100, 1)
        d.add(linea_firma)
        elements.append(d)
        doc.build(elements)
        pdf = buffer.getvalue()
        buffer.close()
        return pdf


def purchase_order_xls_report(order):
    """Construye el libro de Excel de la orden de compra."""
    detalle_index = 0
    quotation = order.quotation
    if quotation is None:
        supplier = order.supplier
    else:
        supplier = order.quotation.supplier
    wb = Workbook()
    ws = wb.active

    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['G'].width = 15
    ws.column_dimensions['I'].width = 15

    ws.merge_cells('B2:I2')
    ws['B2'].font = Font(name='Brush Script MT', size=20, color='0000ff', bold=True)
    ws['B2'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 25
    ws['B2'] = 'Asociación de Bananeros Orgánicos Solidarios  Salitral'
    ws.merge_cells('B3:I3')
    ws['B3'].font = Font(name='Lucida Bright', size=9, color='148918', bold=True, italic=True)
    ws['B3'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B3'] = 'Inscrita en Partida Nº 11003959 de la Superintendencia Nacional de'
    ws.merge_cells('B4:I4')
    ws['B4'].font = Font(name='Lucida Bright', size=9, color='148918', bold=True, italic=True)
    ws['B4'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B4'] = 'los Registros Públicos R.U.C. Nº 20484149748'
    ws.merge_cells('B5:I5')
    ws['B5'].font = Font(name='Lucida Bright', size=9, color='ff0000', bold=True, italic=True)
    ws['B5'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B5'] = 'FLO Certified Banana Producer – FLO ID 2460'
    ws.merge_cells('B6:I6')
    ws['B6'].font = Font(name='Lucida Bright', size=9, color='ff0000', bold=True, italic=True)
    ws['B6'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B6'] = 'Certificate Organic USDA-NOP, EU2092/91, JAS. Registration Nº CU 805878'
    ws.merge_cells('B10:I10')
    ws['B10'].font = Font(name='Calibri', size=12)
    ws['B10'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B10'] = '"AÑO DEL BUEN SERVICIO AL CIUDADANO"'
    ws.merge_cells('D11:F11')
    ws['D11'].font = Font(name='Calibri', size=18, bold=True, underline="single")
    ws['D11'].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[11].height = 25
    ws['D11'] = 'ORDEN DE COMPRA'
    ws['G11'].font = Font(name='Calibri', size=11, bold=True, underline="single")
    ws['G11'].alignment = Alignment(horizontal="center", vertical="center")
    ws['G11'] = 'N°'

    ws['B13'].font = Font(size=11, bold=True)
    ws['B13'] = 'DATOS DEL CLIENTE Y DE LA FACTURA'
    ws['B14'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B14'] = 'RAZÓN SOCIAL'
    ws.merge_cells('C14:G14')
    ws['C14'].alignment = Alignment(horizontal="center")
    ws['C14'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['C14'] = order.supplier.business_name
    ws['H14'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['H14'] = 'FECHA'
    ws['I14'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I14'] = order.date
    ws['B15'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B15'] = 'CONTACTO'
    ws.merge_cells('C15:G15')
    ws['C15'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['C15'].alignment = Alignment(horizontal="center")
    ws['C15'] = ''
    ws['H15'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['H15'] = 'RUC/NIT'
    ws['I15'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I15'] = supplier.tax_id
    ws['B16'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B16'] = 'DIRECCIÓN'
    ws.merge_cells('C16:G16')
    ws['C16'].alignment = Alignment(horizontal="center")
    ws['C16'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['C16'] = supplier.address
    ws['H16'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['H16'] = 'TELÉFONO'
    ws['I16'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    try:
        ws['I16'] = supplier.phone
    except TypeError:
        ws['I16'] = '-'
    ws['B17'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B17'] = 'E-MAIL'
    ws.merge_cells('C17:G17')
    ws['C14'].alignment = Alignment(horizontal="center")
    ws['C17'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['C17'] = supplier.email
    ws['H17'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['H17'] = 'RPM/RPC'
    ws['I17'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I17'] = ''

    ws['B19'].font = Font(size=11, bold=True)
    ws['B19'] = 'DATOS DE ENTREGA'
    ws.merge_cells('B20:C20')
    ws['B20'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B20'] = 'DIRECCIÓN DE ENTREGA'
    ws.merge_cells('D20:F20')
    ws['D20'].alignment = Alignment(horizontal="center")
    ws['D20'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['D20'] = str(company().address())
    ws.merge_cells('G20:H20')
    ws['G20'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G20'] = 'DEPARTAMENTO'
    ws['I20'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I20'] = str(company().department)
    ws.merge_cells('B21:C22')
    ws['B21'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B21'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B21'] = 'PERSONAS AUTORIZADAS'
    ws.merge_cells('D21:F21')
    ws['D21'].alignment = Alignment(horizontal="center")
    ws['D21'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['D21'] = ''
    ws.merge_cells('D22:F22')
    ws['D22'].alignment = Alignment(horizontal="center")
    ws['D22'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['D22'] = ''
    ws.merge_cells('G21:H21')
    ws['G21'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G21'] = 'PROVINCIA'
    ws['I21'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I21'] = str(company().province)
    ws.merge_cells('G22:H22')
    ws['G22'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G22'] = 'DISTRITO'
    ws['I22'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I22'] = str(company().district)
    ws.merge_cells('B23:C24')
    ws['B23'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B23'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B23'] = 'TELÉFONOS'
    ws.merge_cells('D23:F23')
    ws['D23'].alignment = Alignment(horizontal="center")
    ws['D23'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['D23'] = ''
    ws.merge_cells('G23:H23')
    ws['G23'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G23'] = 'COMENTARIO:'
    ws['I23'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I23'] = ''
    ws['B24'].border = Border(bottom=Side(border_style="thin"))
    ws.merge_cells('D24:F24')
    ws['D24'].alignment = Alignment(horizontal="center")
    ws['D24'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['D24'] = ''
    ws.merge_cells('G24:I24')
    ws['D24'].alignment = Alignment(horizontal="center")
    ws['G24'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G24'] = ''
    ws['J24'].border = Border(left=Side(border_style="thin"))
    ws.merge_cells('G25:I25')
    ws['G25'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G25'] = ''
    ws['J25'].border = Border(left=Side(border_style="thin"))
    ws.merge_cells('G26:I26')
    ws['G26'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G26'] = ''
    ws['J26'].border = Border(left=Side(border_style="thin"))

    ws['B28'].font = Font(size=11, bold=True)
    ws['B28'] = 'DATOS DEL PRODUCTO A ADQUIRIR'
    ws.merge_cells('B29:B30')
    ws['B29'].alignment = Alignment(horizontal="center", vertical="center")
    ws['B29'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['B29'] = 'UNIDAD MEDIDA'
    ws.merge_cells('C29:F30')
    ws['C29'].alignment = Alignment(horizontal="center", vertical="center")
    ws['C29'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['C29'] = 'DESCRIPCIÓN DEL PRODUCTO'
    ws.merge_cells('G29:G30')
    ws['G29'].alignment = Alignment(horizontal="center", vertical="center")
    ws['G29'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G29'] = 'CANT.'
    ws['H29'].alignment = Alignment(horizontal="center")
    ws['H29'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['H29'] = 'PRECIO'
    ws['I29'].alignment = Alignment(horizontal="center")
    ws['I29'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I29'] = 'PRECIO'
    ws['H30'].alignment = Alignment(horizontal="center")
    ws['H30'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['H30'] = 'UNITARIO'
    ws['I30'].alignment = Alignment(horizontal="center")
    ws['I30'].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I30'] = 'TOTAL'

    for item in PurchaseOrderDetail.objects.filter(order=order):
        fila = 31 + detalle_index
        ws['B' + str(fila)].alignment = Alignment(horizontal="center")
        ws['B' + str(fila)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                            top=Side(border_style="thin"), bottom=Side(border_style="thin"))
        ws['B' + str(fila)] = item.product.unit_of_measure.description
        ws.merge_cells('C' + str(fila) + ':F' + str(fila))
        ws['C' + str(fila)].alignment = Alignment(horizontal="center")
        ws['C' + str(fila)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                            top=Side(border_style="thin"), bottom=Side(border_style="thin"))
        ws['C' + str(fila)] = item.product.description
        ws['G' + str(fila)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                            top=Side(border_style="thin"), bottom=Side(border_style="thin"))
        ws['G' + str(fila)] = item.quantity
        ws['H' + str(fila)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                            top=Side(border_style="thin"), bottom=Side(border_style="thin"))
        ws['H' + str(fila)] = item.price
        ws['I' + str(fila)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                            top=Side(border_style="thin"), bottom=Side(border_style="thin"))
        ws['I' + str(fila)] = item.quantity * item.price
        detalle_index += 1

    fila_total = 31 + detalle_index
    ws.merge_cells('G' + str(fila_total) + ':H' + str(fila_total))
    ws['G' + str(fila_total)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_total)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G' + str(fila_total)] = 'SUBTOTAL'
    ws['I' + str(fila_total)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                              top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I' + str(fila_total)] = order.subtotal
    ws.merge_cells('G' + str(fila_total + 1) + ':H' + str(fila_total + 1))
    ws['G' + str(fila_total + 1)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_total + 1)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                                  top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G' + str(fila_total + 1)] = 'IMPUESTO 18% IGV'
    ws['I' + str(fila_total + 1)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                                  top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I' + str(fila_total + 1)] = order.impuesto
    ws.merge_cells('G' + str(fila_total + 2) + ':H' + str(fila_total + 2))
    ws['G' + str(fila_total + 2)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_total + 2)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                                  top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['G' + str(fila_total + 2)] = 'TOTAL'
    ws['I' + str(fila_total + 2)].border = Border(left=Side(border_style="thin"), right=Side(border_style="thin"),
                                                  top=Side(border_style="thin"), bottom=Side(border_style="thin"))
    ws['I' + str(fila_total + 2)] = order.total

    ws['B' + str(fila_total + 7)] = 'SON:'
    ws.merge_cells('C' + str(fila_total + 7) + ':I' + str(fila_total + 7))
    ws['C' + str(fila_total + 7)].font = Font(underline="single")
    ws['C' + str(fila_total + 7)].alignment = Alignment(horizontal="center")
    ws['C' + str(fila_total + 7)] = order.total_in_words

    fila_pago = fila_total + 9
    ws.merge_cells('E' + str(fila_pago) + ':F' + str(fila_pago))
    ws['E' + str(fila_pago)] = 'FORMA DE PAGO'

    ws.merge_cells('G' + str(fila_pago) + ':I' + str(fila_pago))
    ws['G' + str(fila_pago)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_pago)].border = Border(bottom=Side(border_style="thin"))
    ws['G' + str(fila_pago)] = order.payment_method.description
    ws.merge_cells('E' + str(fila_pago + 1) + ':F' + str(fila_pago + 1))
    ws['E' + str(fila_pago + 1)] = 'BANCO'
    ws.merge_cells('G' + str(fila_pago + 1) + ':I' + str(fila_pago + 1))
    ws['G' + str(fila_pago + 1)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_pago + 1)].border = Border(bottom=Side(border_style="thin"))
    ws['G' + str(fila_pago + 1)] = ''

    ws.merge_cells('E' + str(fila_pago + 2) + ':F' + str(fila_pago + 2))
    ws['E' + str(fila_pago + 2)] = 'NÚMERO DE CTA:'
    ws.merge_cells('G' + str(fila_pago + 2) + ':I' + str(fila_pago + 2))
    ws['G' + str(fila_pago + 2)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_pago + 2)].border = Border(bottom=Side(border_style="thin"))
    ws['G' + str(fila_pago + 2)] = ''

    ws.merge_cells('E' + str(fila_pago + 3) + ':F' + str(fila_pago + 3))
    ws['E' + str(fila_pago + 3)] = 'TIPO DE CUENTA'
    ws.merge_cells('G' + str(fila_pago + 3) + ':I' + str(fila_pago + 3))
    ws['G' + str(fila_pago + 3)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_pago + 3)].border = Border(bottom=Side(border_style="thin"))
    ws['G' + str(fila_pago + 3)] = ''

    ws.merge_cells('E' + str(fila_pago + 4) + ':F' + str(fila_pago + 4))
    ws['E' + str(fila_pago + 4)] = 'MONEDA'
    ws.merge_cells('G' + str(fila_pago + 4) + ':I' + str(fila_pago + 4))
    ws['G' + str(fila_pago + 4)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_pago + 4)].border = Border(bottom=Side(border_style="thin"))
    ws['G' + str(fila_pago + 4)] = ''

    ws.merge_cells('E' + str(fila_pago + 5) + ':F' + str(fila_pago + 5))
    ws['E' + str(fila_pago + 5)] = 'CCI'
    ws.merge_cells('G' + str(fila_pago + 5) + ':I' + str(fila_pago + 5))
    ws['G' + str(fila_pago + 5)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_pago + 5)].border = Border(bottom=Side(border_style="thin"))
    ws['G' + str(fila_pago + 5)] = ''

    fila_firmas = fila_pago + 11
    ws.merge_cells('C' + str(fila_firmas) + ':E' + str(fila_firmas))
    ws['C' + str(fila_firmas)].alignment = Alignment(horizontal="center")
    ws['C' + str(fila_firmas)].border = Border(top=Side(border_style="thin"))
    ws['C' + str(fila_firmas)] = 'RESPONSABLE'
    ws.merge_cells('G' + str(fila_firmas) + ':I' + str(fila_firmas))
    ws['G' + str(fila_firmas)].alignment = Alignment(horizontal="center")
    ws['G' + str(fila_firmas)].border = Border(top=Side(border_style="thin"))
    ws['G' + str(fila_firmas)] = 'APROBADO POR'
    return wb


class QuotationRequestPdf(object):

    def header(self, pdf, quotation):
        try:
            archivo_imagen = os.path.join(settings.MEDIA_ROOT, str(company().logo))
            pdf.drawImage(archivo_imagen, 20, 750, 120, 90, preserveAspectRatio=True)
        except Exception:
            pdf.drawString(20, 800, 'LOGO')
        pdf.setFont("Times-Roman", 14)
        encabezado = [[u"SOLICITUD DE COTIZACIÓN"]]
        header_table = Table(encabezado, colWidths=[8 * cm])
        header_table.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (0, 0), 'CENTER'),
                ('GRID', (0, 0), (1, 0), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
            ]
        ))
        header_table.wrapOn(pdf, 800, 600)
        header_table.drawOn(pdf, 200, 800)
        pdf.drawString(270, 780, u"N°" + quotation.code)
        pdf.setFont("Times-Roman", 10)
        pdf.drawString(40, 750, u"SEÑOR(ES): " + quotation.supplier.business_name)
        pdf.drawString(440, 750, u"R.U.C.: " + quotation.supplier.tax_id)
        address = quotation.supplier.address
        if len(address) > 60:
            pdf.drawString(40, 730, u"DIRECCIÓN: " + address[0:60])
            pdf.drawString(105, 720, address[60:])
        else:
            pdf.drawString(40, 730, u"DIRECCIÓN: " + address)
        try:
            pdf.drawString(440, 730, u"TELÉFONO: " + quotation.supplier.phone)
        except TypeError:
            pdf.drawString(440, 730, u"TELÉFONO: -")
        pdf.drawString(40, 710, u"FECHA: " + quotation.date.strftime('%d/%m/%Y'))

    def detail(self, pdf, y, quotation):
        encabezados = ('Nro', 'Descripción', 'Unidad', 'Cantidad')
        detalles = quotation.details.all()
        lista_detalles = []
        for detail in detalles:
            tupla_producto = (detail.line_number, detail.requirement_detail.product.description,
                              detail.requirement_detail.product.unit_of_measure.description, detail.quantity)
            lista_detalles.append(tupla_producto)
        adicionales = [('', '', '', '')] * (15 - len(detalles))
        tabla_detalle = Table([encabezados] + lista_detalles + adicionales,
                              colWidths=[1 * cm, 13.5 * cm, 1.5 * cm, 2 * cm])
        tabla_detalle.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (3, 0), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 7),
                ('ALIGN', (3, 1), (-1, -1), 'LEFT'),
            ]
        ))
        tabla_detalle.wrapOn(pdf, 800, 600)
        tabla_detalle.drawOn(pdf, 40, y + 80)

    def notes_box(self, pdf, y, quotation):
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 8
        p.fontName = "Times-Roman"
        obs = Paragraph("OBSERVACIONES: " + quotation.notes, p)
        notes = [[obs]]
        notes_table = Table(notes, colWidths=[18 * cm], rowHeights=1.8 * cm)
        notes_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (0, 2), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]
        ))
        notes_table.wrapOn(pdf, 800, 600)
        notes_table.drawOn(pdf, 40, y + 20)

    def render(self, quotation):
        """Devuelve el PDF ya generado."""
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer)
        self.header(pdf, quotation)
        y = 300
        self.detail(pdf, y, quotation)
        self.notes_box(pdf, y, quotation)
        ''''
        self.firmas(pdf, y, quotation)'''
        pdf.showPage()
        pdf.save()
        contenido = buffer.getvalue()
        buffer.close()
        return contenido



class ServiceConformityMemoPdf(object):

    def get_position(self, office, conformity):
        try:
            position = Position.objects.get(office=office,
                                        is_leadership=True,
                                        start_date__lte=conformity.date,
                                        end_date=None)
        except Position.DoesNotExist:
            position = Position.objects.get(office=office,
                                        is_leadership=True,
                                        start_date__lte=conformity.date,
                                        end_date__gte=conformity.date)
        return position

    def superior_position(self, office, conformity):
        try:
            superior_position = Position.objects.get(office=office,
                                                 is_leadership=True,
                                                 start_date__lte=conformity.date,
                                                 end_date=None)
        except Position.DoesNotExist:
            superior_position = Position.objects.get(office=office,
                                                 is_leadership=True,
                                                 start_date__lte=conformity.date,
                                                 end_date__gte=conformity.date)
        return superior_position

    def header(self, pdf, conformity):
        try:
            archivo_imagen = os.path.join(settings.MEDIA_ROOT, str(company().logo))
            pdf.drawImage(archivo_imagen, 40, 750, 100, 70, preserveAspectRatio=True)
        except Exception:
            pdf.drawString(40, 750, 'LOGO')
        pdf.setFont("Times-Roman", 14)
        pdf.drawString(130, 750, u"MEMORANDO DE CONFORMIDAD DEL SERVICIO")
        pdf.setFont("Times-Roman", 13)
        pdf.drawString(250, 730, u"N°" + conformity.code)
        pdf.setFont("Times-Roman", 10)
        pdf.drawString(430, 780, company().district + " " + conformity.date.strftime('%d de %b de %Y'))
        pdf.drawString(475, 710, conformity.service_order.code)
        requirement = conformity.service_order.quotation.requirement
        gerencia_inmediata = requirement.office.management
        requester = requirement.requester
        puesto_solicitante = requester.position
        if puesto_solicitante is None:
            puesto_solicitante = self.get_position(requirement.office, conformity)
        puesto_jefe_inmediato = self.superior_position(requirement.office, conformity)
        jefe_inmediato = puesto_jefe_inmediato.worker
        y = 690
        if puesto_solicitante.office.code == 'GGEN':
            puesto_gerente = self.get_position(configuration().administration, conformity)
        elif puesto_solicitante.office.code == 'GOPE' and not puesto_solicitante.is_leadership:
            puesto_gerente = self.get_position(requirement.office, conformity)
        else:
            puesto_gerente = self.get_position(gerencia_inmediata, conformity)
        gerente = puesto_gerente.worker
        if puesto_gerente.pk == puesto_jefe_inmediato.pk or puesto_jefe_inmediato.pk == puesto_solicitante.pk:
            pdf.drawString(50, y, u"A           :    " + gerente.full_name())
            y = y - 20
            pdf.drawString(50, y, u"                   " + puesto_gerente.name)
            y = y - 20
            pdf.drawString(50, y, u"DE        :     " + puesto_solicitante.worker.full_name())
            y = y - 20
            pdf.drawString(50, y, u"                   " + puesto_solicitante.name)
            y = y - 50
        else:
            pdf.drawString(50, y, u"A           :    " + gerente.full_name())
            y = y - 20
            pdf.drawString(50, y, u"                   " + puesto_gerente.name)
            y = y - 20
            pdf.drawString(50, y, u"                   " + jefe_inmediato.full_name())
            y = y - 20
            pdf.drawString(50, y, u"                   " + puesto_jefe_inmediato.name)
            y = y - 20
            pdf.drawString(50, y, u"DE        :     " + puesto_solicitante.worker.full_name())
            y = y - 20
            pdf.drawString(50, y, u"                   " + puesto_solicitante.name)
            y = y - 30
        estilo_parrafo = ParagraphStyle('parrafos')
        estilo_parrafo.alignment = TA_JUSTIFY
        estilo_parrafo.fontSize = 10
        estilo_parrafo.fontName = "Times-Roman"
        cadena_parrafo = u"""Mediante el presente comunico a Ud. que el servicio requerido con REQ DE BIENES Y SERV. N° %s, 
        ha sido concluido a satisfacción, según %s, lo que comunicamos para que proceda al pago del servicio correspondiente que 
        se detalla como sigue: """ % (requirement.code, conformity.supporting_document)
        p1 = Paragraph(cadena_parrafo, estilo_parrafo)
        p1.wrapOn(pdf, 500, y - 20)
        p1.drawOn(pdf, 40, y - 20)

    def detail(self, pdf, y, conformity):
        encabezados = ('ITEM', 'DETALLE DEL SERVICIO')
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 9
        p.fontName = "Times-Roman"
        detalles = []
        cont = 0
        for detail in ServiceConformityDetail.objects.filter(conformity=conformity):
            description = detail.service_order_detail.quotation_detail.requirement_detail.product.description + '-' + detail.service_order_detail.quotation_detail.requirement_detail.use
            if len(description) > 58:
                cont = cont + 1
            detalles.append((detail.line_number, Paragraph(description, p)))
        adicionales = [('', '')] * (8 - cont - len(detalles))
        detalle_orden = Table([encabezados] + detalles + adicionales, colWidths=[0.8 * cm, 17 * cm])
        detalle_orden.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (1, 0), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                # ('LINEBELOW', (0,1), (5,-1), 0, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (1, 1), (-1, -1), 'RIGHT'),
            ]
        ))
        detalle_orden.wrapOn(pdf, 800, 600)
        detalle_orden.drawOn(pdf, 40, y + 75)

    def signature(self, pdf, x_texto, y_texto, texto, x_ini_linea, x_fin_linea, y_linea):
        pdf.drawString(x_texto, y_texto, texto)
        pdf.line(x_ini_linea, y_linea, x_fin_linea, y_linea)

    def render(self, conformity):
        """Devuelve el PDF ya generado."""
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer)
        self.header(pdf, conformity)
        y = 300
        self.detail(pdf, y, conformity)
        pdf.setFont("Times-Roman", 8)
        self.signature(pdf, 170, y - 50, "GERENCIA", 120, 265, y - 40)
        self.signature(pdf, 330, y - 50, "CONFORMIDAD DEL SOLICITANTE", 320, 470, y - 40)
        self.signature(pdf, 130, y - 150, "CONFORMIDAD JEFE INMEDIATO", 120, 265, y - 140)
        self.signature(pdf, 350, y - 150, "UNIDAD DE LOGÍSTICA", 320, 470, y - 140)
        pdf.drawCentredString(300, y - 280, company().address())
        pdf.showPage()
        pdf.save()
        contenido = buffer.getvalue()
        buffer.close()
        return contenido



class ServiceOrderPdf(object):

    def header(self, pdf, order):
        try:
            archivo_imagen = os.path.join(settings.MEDIA_ROOT, str(company().logo))
            pdf.drawImage(archivo_imagen, 40, 750, 120, 90, preserveAspectRatio=True)
        except Exception:
            pdf.drawString(40, 800, 'LOGO')
        pdf.setFont("Times-Roman", 14)
        pdf.drawString(230, 800, u"ORDEN DE SERVICIOS")
        pdf.setFont("Times-Roman", 11)
        pdf.drawString(455, 800, u"R.U.C. " + company().tax_id)
        pdf.setFont("Times-Roman", 13)
        pdf.drawString(250, 780, u"N°" + order.code)
        pdf.setFont("Times-Roman", 10)
        pdf.drawString(430, 780, company().district + " " + order.date.strftime('%d de %b de %Y'))
        pdf.setFont("Times-Roman", 10)
        quotation = order.quotation
        if quotation is None:
            supplier = order.supplier
        else:
            supplier = order.quotation.supplier
        pdf.drawString(40, 750, u"SEÑOR(ES): " + supplier.business_name)
        pdf.drawString(440, 750, u"R.U.C.: " + supplier.tax_id)
        address = supplier.address
        if len(address) > 60:
            pdf.drawString(40, 730, u"DIRECCIÓN: " + address[0:60])
            pdf.drawString(105, 720, address[60:])
        else:
            pdf.drawString(40, 730, u"DIRECCIÓN: " + address)
        try:
            pdf.drawString(440, 730, u"TELÉFONO: " + supplier.phone)
        except TypeError:
            pdf.drawString(440, 730, u"TELÉFONO: -")
        try:
            pdf.drawString(40, 710,
                           u"REFERENCIA: " + order.quotation.requirement.code + " - " + order.quotation.requirement.office.name)
        except (ObjectDoesNotExist, AttributeError):
            pdf.drawString(40, 710, u"REFERENCIA: " + order.report_name)

        pdf.drawString(40, 690, u"PROCESO: " + order.process)
        pdf.setFont("Times-Roman", 8)
        pdf.drawString(40, 670, u"Sírvase remitirnos según especificaciones que detallamos lo next: ")

    def detail(self, pdf, y, order):
        encabezados = ('Item', 'Cantidad', u'Descripción', 'Precio', 'Total')
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 9
        p.fontName = "Times-Roman"
        detalles = []
        cont = 0

        for detail in ServiceOrderDetail.objects.filter(order=order):
            try:
                description = detail.quotation_detail.requirement_detail.product.description
                if len(description) > 58:
                    cont = cont + 1
                detalles.append(
                    (detail.line_number, detail.quantity, Paragraph(description, p), detail.price, detail.amount))
            except (ObjectDoesNotExist, AttributeError):
                description = detail.product.description
                if len(description) > 58:
                    cont = cont + 1
                detalles.append(
                    (detail.line_number, detail.quantity, Paragraph(description, p), detail.price, detail.amount))

        # detalles = [(detail.line_number, detail.quantity, Paragraph(detail.servicio.description+'-'+detail.description,p), detail.price,detail.amount) for detail in ServiceOrderDetail.objects.filter(order=order)]
        adicionales = [('', '', '', '', '')] * (15 - cont - len(detalles))
        detalle_orden = Table([encabezados] + detalles + adicionales,
                              colWidths=[0.8 * cm, 1.9 * cm, 11.3 * cm, 2 * cm, 2.5 * cm])
        detalle_orden.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (4, 0), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                # ('LINEBELOW', (0,1), (5,-1), 0, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (4, 1), (-1, -1), 'RIGHT'),
            ]
        ))
        detalle_orden.wrapOn(pdf, 800, 600)
        detalle_orden.drawOn(pdf, 40, y + 75)
        # Letras
        total_in_words = [("SON: " + order.total_in_words, '')]
        total_in_words_table = Table(total_in_words, colWidths=[16 * cm, 2.5 * cm])
        total_in_words_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (1, 0), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        total_in_words_table.wrapOn(pdf, 800, 600)
        total_in_words_table.drawOn(pdf, 40, y + 55)

    def others(self, pdf, y, order):
        encabezados_otros = ('LUGAR DE ENTREGA', 'PLAZO DE ENTREGA', 'FORMA DE PAGO')
        others = [('', u" DÍAS", "")]
        others_table = Table([encabezados_otros] + others, colWidths=[6 * cm, 3.5 * cm, 4.5 * cm],
                            rowHeights=[0.6 * cm, 1 * cm])
        others_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (2, 1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        others_table.wrapOn(pdf, 800, 600)
        others_table.drawOn(pdf, 40, y + 5)

    def total_box(self, pdf, y, order):
        pdf.drawString(445, y + 40, u"SUB-TOTAL: ")
        pdf.drawString(445, y + 20, u"IGV: ")
        pdf.drawString(445, y, u"TOTAL: ")
        total = [[order.subtotal], [str(order.impuesto)], [str(order.total)]]
        total_table = Table(total, colWidths=[2.5 * cm])
        total_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (0, 2), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
            ]
        ))
        total_table.wrapOn(pdf, 800, 600)
        total_table.drawOn(pdf, 495, y - 2)

    def notes_box(self, pdf, y, order):
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 10
        p.fontName = "Times-Roman"
        obs = Paragraph("Observaciones: " + order.notes, p)
        notes = [[obs]]
        notes_table = Table(notes, colWidths=[18.50 * cm], rowHeights=1.8 * cm)
        notes_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (0, 2), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]
        ))
        notes_table.wrapOn(pdf, 800, 600)
        notes_table.drawOn(pdf, 40, y - 58)

    def budget_impact(self, pdf):
        y = 320
        pdf.drawString(40, y - 90, u"HOJA DE AFECTACIÓN PRESUPUESTAL:")
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 8
        p.fontName = "Times-Roman"
        lista = ListFlowable([
            Paragraph("""Consignar el número de la presente Orden de Compra en su Guía de Remisión y Factura. 
                          Facturar a nombre de """ + force_str(company().business_name), p),
            Paragraph("El " + force_str(company().business_name) + """, se reserva el derecho de devolver 
                          la mercaderia, sino se ajusta a las especificaciones requeridas, asimismo de anular la presente 
                          Orden de Compra.""", p),
            Paragraph("""El pago de toda factura se hará de acuerdo a las condiciones establecidas.""", p)
        ], bulletType='1'
        )
        pdf.drawString(330, y - 150, "FIRMA: ")
        pdf.line(370, y - 150, 560, y - 150)
        pdf.drawString(330, y - 170, "NOMBRE: ")
        pdf.line(370, y - 170, 560, y - 170)
        pdf.drawString(330, y - 190, "DNI: ")
        pdf.line(370, y - 190, 560, y - 190)
        pdf.setFont("Times-Roman", 6)
        pdf.drawString(525, y - 127, "FECHA")
        afectacion = [[(Paragraph("IMPORTANTE:", p), lista), "RECIBIDO POR:"]]
        tabla_afectacion = Table(afectacion, colWidths=[10 * cm, 8.50 * cm])
        tabla_afectacion.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (1, 0), 1, colors.black),
                ('VALIGN', (1, 0), (1, 0), 'TOP'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        tabla_afectacion.wrapOn(pdf, 800, 600)
        tabla_afectacion.drawOn(pdf, 40, y - 200)
        date = [[' ', ' ', ' ']]
        tabla_date = Table(date, colWidths=[0.6 * cm, 0.6 * cm, 0.6 * cm], rowHeights=0.6 * cm)
        tabla_date.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 5),
            ]
        ))
        tabla_date.wrapOn(pdf, 800, 600)
        tabla_date.drawOn(pdf, 510, y - 120)

    def render(self, order):
        """Devuelve el PDF ya generado."""
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer)
        self.header(pdf, order)
        y = 300
        self.detail(pdf, y, order)
        self.others(pdf, y, order)
        self.total_box(pdf, y, order)
        self.notes_box(pdf, y, order)
        self.budget_impact(pdf)
        pdf.setFont("Times-Roman", 8)
        pdf.drawString(115, y - 250, "Elaborado por")
        pdf.drawString(430, y - 250, "Autorizado por")
        pdf.line(70, y - 240, 200, y - 240)
        pdf.line(390, y - 240, 520, y - 240)
        pdf.drawCentredString(300, y - 280, company().address())
        pdf.showPage()
        pdf.save()
        contenido = buffer.getvalue()
        buffer.close()
        return contenido



class PurchaseOrderPdf(object):

    def header(self, pdf, order):
        try:
            archivo_imagen = os.path.join(settings.MEDIA_ROOT, str(company().logo))
            pdf.drawImage(archivo_imagen, 40, 750, 100, 90, mask='auto', preserveAspectRatio=True)
        except Exception:
            pdf.drawString(40, 800, 'LOGO')
        pdf.setFont("Times-Roman", 14)
        pdf.drawString(230, 800, u"ORDEN DE COMPRA")
        pdf.setFont("Times-Roman", 11)
        pdf.drawString(455, 800, u"R.U.C. " + company().tax_id)
        pdf.setFont("Times-Roman", 13)
        pdf.drawString(250, 780, u"N° " + order.code)
        pdf.setFont("Times-Roman", 10)
        pdf.drawString(430, 780, company().district + " " + order.date.strftime(
            '%d de %b de %Y'))  # order.date.strftime('%d de %B de %Y')
        pdf.setFont("Times-Roman", 10)
        quotation = order.quotation
        if quotation is None:
            supplier = order.supplier
        else:
            supplier = order.quotation.supplier
        pdf.drawString(40, 750, u"SEÑOR(ES): " + supplier.business_name)
        pdf.drawString(440, 750, u"R.U.C.: " + supplier.tax_id)
        address = supplier.address
        if len(address) > 60:
            pdf.drawString(40, 730, u"DIRECCIÓN: " + address[0:60])
            pdf.drawString(105, 720, address[60:])
        else:
            pdf.drawString(40, 730, u"DIRECCIÓN: " + address)
        try:
            pdf.drawString(440, 730, u"TELÉFONO: " + supplier.phone)
        except TypeError:
            pdf.drawString(440, 730, u"TELÉFONO: -")
        try:
            pdf.drawString(40, 710,
                           u"REFERENCIA: " + order.quotation.requirement.code + " - " + order.quotation.requirement.office.name)
        except (ObjectDoesNotExist, AttributeError):
            pdf.drawString(40, 710, u"REFERENCIA: -")
        pdf.drawString(40, 690, u"PROCESO: -")
        pdf.setFont("Times-Roman", 8)
        pdf.drawString(40, 670, u"Sírvase remitirnos según especificaciones que detallamos lo next: ")

    def detail(self, pdf, y, order):
        encabezados = ('Item', 'Cantidad', 'Unidad', u'Descripción', 'Precio', 'Total')
        try:
            detalles = [(detail.line_number, detail.quantity,
                         detail.quotation_detail.requirement_detail.product.unit_of_measure.description,
                         detail.quotation_detail.requirement_detail.product.description, detail.price,
                         round(detail.amount, 5)) for detail in PurchaseOrderDetail.objects.filter(order=order)]
        except (ObjectDoesNotExist, AttributeError):
            detalles = [(detail.line_number, detail.quantity, detail.product.unit_of_measure.description,
                         detail.product.description, detail.price, round(detail.price, 5)) for detail in
                        PurchaseOrderDetail.objects.filter(order=order)]
        adicionales = [('', '', '', '', '', '')] * (15 - len(detalles))
        detalle_orden = Table([encabezados] + detalles + adicionales,
                              colWidths=[0.8 * cm, 1.9 * cm, 2 * cm, 9.3 * cm, 2 * cm, 2.5 * cm])
        detalle_orden.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (5, 0), 'CENTER'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                # ('LINEBELOW', (0,1), (5,-1), 0, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (4, 1), (-1, -1), 'RIGHT'),
            ]
        ))
        detalle_orden.wrapOn(pdf, 800, 600)
        detalle_orden.drawOn(pdf, 40, y + 75)
        # Letras
        total_in_words = [("SON: " + order.total_in_words, '')]
        total_in_words_table = Table(total_in_words, colWidths=[16 * cm, 2.5 * cm])
        total_in_words_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (1, 0), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        total_in_words_table.wrapOn(pdf, 800, 600)
        total_in_words_table.drawOn(pdf, 40, y + 55)

    def others(self, pdf, y, order):
        encabezados_otros = ('LUGAR DE ENTREGA', 'PLAZO DE ENTREGA', 'FORMA DE PAGO')
        others = [(company().address(), u"INMEDIATA", order.payment_method.description)]
        others_table = Table([encabezados_otros] + others, colWidths=[6 * cm, 3.5 * cm, 4.5 * cm],
                            rowHeights=[0.6 * cm, 1 * cm])
        others_table.setStyle(TableStyle(
            [
                ('ALIGN', (0, 0), (2, 0), 'CENTER'),
                ('GRID', (0, 0), (2, 1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        others_table.wrapOn(pdf, 800, 600)
        others_table.drawOn(pdf, 40, y + 5)

    def total_box(self, pdf, y, order):
        pdf.drawString(445, y + 40, u"SUB-TOTAL: ")
        pdf.drawString(445, y + 20, u"IGV: ")
        pdf.drawString(445, y, u"TOTAL: S/")
        total = [[round(order.subtotal, 2)], [round(order.impuesto, 2)], [round(order.total, 2)]]
        total_table = Table(total, colWidths=[2.5 * cm])
        total_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (0, 2), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
            ]
        ))
        total_table.wrapOn(pdf, 800, 600)
        total_table.drawOn(pdf, 495, y - 2)

    def notes_box(self, pdf, y, order):
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 10
        p.fontName = "Times-Roman"
        obs = Paragraph("Observaciones: " + order.notes, p)
        notes = [[obs]]
        notes_table = Table(notes, colWidths=[18.50 * cm], rowHeights=1.8 * cm)
        notes_table.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (0, 2), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]
        ))
        notes_table.wrapOn(pdf, 800, 600)
        notes_table.drawOn(pdf, 40, y - 58)

    def budget_impact(self, pdf):
        y = 320
        pdf.drawString(40, y - 90, u"HOJA DE AFECTACIÓN PRESUPUESTAL:")
        p = ParagraphStyle('parrafos')
        p.alignment = TA_JUSTIFY
        p.fontSize = 8
        p.fontName = "Times-Roman"
        lista = ListFlowable([
            Paragraph("""Consignar el número de la presente Orden de Compra en su Guía de Remisión y Factura. 
                          Facturar a nombre de """ + force_str(company().business_name), p),
            Paragraph("El " + force_str(company().business_name) + """, se reserva el derecho de devolver 
                          la mercaderia, sino se ajusta a las especificaciones requeridas, asimismo de anular la presente 
                          Orden de Compra.""", p),
            Paragraph("""El pago de toda factura se hará de acuerdo a las condiciones establecidas.""", p)
        ], bulletType='1'
        )
        pdf.drawString(330, y - 150, "FIRMA: ")
        pdf.line(370, y - 150, 560, y - 150)
        pdf.drawString(330, y - 170, "NOMBRE: ")
        pdf.line(370, y - 170, 560, y - 170)
        pdf.drawString(330, y - 190, "DNI: ")
        pdf.line(370, y - 190, 560, y - 190)
        pdf.setFont("Times-Roman", 6)
        pdf.drawString(525, y - 127, "FECHA")
        afectacion = [[(Paragraph("IMPORTANTE:", p), lista), "RECIBIDO POR:"]]
        tabla_afectacion = Table(afectacion, colWidths=[10 * cm, 8.50 * cm])
        tabla_afectacion.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (1, 0), 1, colors.black),
                ('VALIGN', (1, 0), (1, 0), 'TOP'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
            ]
        ))
        tabla_afectacion.wrapOn(pdf, 800, 600)
        tabla_afectacion.drawOn(pdf, 40, y - 200)
        date = [[' ', ' ', ' ']]
        tabla_date = Table(date, colWidths=[0.6 * cm, 0.6 * cm, 0.6 * cm], rowHeights=0.6 * cm)
        tabla_date.setStyle(TableStyle(
            [
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 0), (-1, -1), 5),
            ]
        ))
        tabla_date.wrapOn(pdf, 800, 600)
        tabla_date.drawOn(pdf, 510, y - 120)

    def render(self, order):
        """Devuelve el PDF ya generado."""
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer)
        self.header(pdf, order)
        y = 300
        self.detail(pdf, y, order)
        self.others(pdf, y, order)
        self.total_box(pdf, y, order)
        self.notes_box(pdf, y, order)
        self.budget_impact(pdf)
        pdf.setFont("Times-Roman", 8)
        pdf.drawString(115, y - 250, "Elaborado por")
        pdf.drawString(430, y - 250, "Autorizado por")
        pdf.line(70, y - 240, 200, y - 240)
        pdf.line(390, y - 240, 520, y - 240)
        pdf.drawCentredString(300, y - 280, company().address())
        pdf.showPage()
        pdf.save()
        contenido = buffer.getvalue()
        buffer.close()
        return contenido
