# -*- coding: utf-8 -*-
import logging

from django.shortcuts import render
from administration.forms import OfficeForm, WorkerForm, PositionForm, PositionUpdateForm,\
    ProfessionForm, ApprovalLevelForm, ProducerForm
from warehouse.models import MovementType
from accounting.forms import UploadForm
from tambox.views import CsvImportMixin, AjaxOnlyMixin
from django.urls import reverse_lazy
from django.views.generic.edit import FormView, UpdateView, CreateView
from django.views.generic.list import ListView
from administration.models import Office, Worker, Position, Profession,\
    ApprovalLevel, Producer
from django.views.generic.base import View, TemplateView
from django.views.generic.detail import DetailView
from django.urls import reverse
from django.contrib.auth.models import User
from openpyxl import Workbook
from django.http import HttpResponse
import datetime
from security.permisos import requires
from django.utils.decorators import method_decorator
import simplejson
import json
from django.db.models import Q

logger = logging.getLogger(__name__)


class Dashboard(View):

    def get(self, request, *args, **kwargs):
        lista_notificaciones = []
        worker_count = Worker.objects.all().count()
        position_count = Position.objects.all().count()
        profession_count = Profession.objects.all().count()
        office, creada = Office.objects.get_or_create(code='GGEN',
                                                       defaults={'name': 'GERENCIA GENERAL',
                                                                 'is_management': True})
        if creada:
            lista_notificaciones.append("Se ha creado la oficina de GERENCIA GENERAL")
        if worker_count == 0:
            lista_notificaciones.append("No se ha registrado ningún trabajador")
        if position_count == 0:
            lista_notificaciones.append("No se ha registrado ningún puesto")
        if profession_count == 0:
            lista_notificaciones.append("No se ha registrado ninguna profesión")
        logistics_level, creada = ApprovalLevel.objects.get_or_create(description="LOGISTICA")
        _, creado = ApprovalLevel.objects.get_or_create(description="USUARIO",
                                                         defaults={'superior_level': logistics_level})
        if creada or creado:
            lista_notificaciones.append("Se han creado los niveles de aprobación básicos")
        context = {'notificaciones': lista_notificaciones}
        return render(request, 'administration/administration_dashboard.html', context)


class ReceiverDniSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('dni', 'movement_type')
    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            dni = request.GET['dni']
            movement_type = MovementType.objects.get(pk=request.GET['movement_type'])
            if movement_type.is_sale:
                receiver = Producer.objects.get(dni=dni)
            else:
                receiver = Worker.objects.get(dni=dni)
            receptor_json = {}
            receptor_json['dni'] = receiver.dni
            receptor_json['full_name'] = str(receiver.full_name())
            data = simplejson.dumps(receptor_json)
            return HttpResponse(data, 'application/json')


class ReceiverNameSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('name', 'movement_type')
    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            name = request.GET['name']
            movement_type = MovementType.objects.get(pk=request.GET['movement_type'])
            if movement_type.is_sale:
                receptores = Producer.objects.filter(last_name__icontains=name)[:20]
            else:
                receptores = Worker.objects.filter(
                    Q(last_name__icontains=name) | Q(
                        first_name__icontains=name))[:20]
            lista_receptores = []
            for receiver in receptores:
                receptor_json = {}
                receptor_json['label'] = str(receiver.full_name())
                receptor_json['dni'] = receiver.dni
                lista_receptores.append(receptor_json)
            data = json.dumps(lista_receptores)
            return HttpResponse(data, 'application/json')


class OfficeImport(CsvImportMixin, FormView):
    template_name = 'administration/office_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('administration:office_list')

    def process_row(self, row):
        Office.objects.get_or_create(code=row[0],
                                      defaults={
                                          'name': row[1],
                                          'dependency': Office.objects.get(code=row[2])},
                                      )


class ProducerImport(CsvImportMixin, FormView):
    template_name = 'administration/producer_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('administration:producer_list')

    def process_row(self, row):
        dni = row[0]
        if dni != "":
            try:
                Producer.objects.get_or_create(dni=dni,
                                                defaults={'last_name': (row[1] + ' ' + row[2]).upper(),
                                                          'first_name': row[3].upper()})
            except Exception:
                logger.warning("No se pudo importar el productor con DNI %s", dni, exc_info=True)


class WorkerImport(CsvImportMixin, FormView):
    template_name = 'administration/worker_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('administration:worker_list')

    def process_row(self, row):
        sheet_user = row[0]
        if sheet_user != "":
            user, creado = User.objects.get_or_create(username=sheet_user,
                                                         defaults={'email': row[5]}, )
            if creado:
                user.set_unusable_password()
                user.save()
                Worker.objects.get_or_create(user=user,
                                                 defaults={'dni': row[1].strip(),
                                                           'last_name': (row[2] + ' ' + row[3]).strip(),
                                                           'first_name': row[4]})
        else:
            Worker.objects.get_or_create(dni=row[1].strip(),
                                             defaults={'last_name': (row[2] + ' ' + row[3]).strip(),
                                                       'first_name': row[4]})


class PositionImport(CsvImportMixin, FormView):
    template_name = 'administration/position_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('administration:position_list')

    def process_row(self, row):
        date = datetime.date(int(row[3][6:]), int(row[3][3:5]), int(row[3][0:2]))
        try:
            Position.objects.get_or_create(name=row[0],
                                         defaults={'office': Office.objects.get(code=row[1].strip()),
                                                   'worker': Worker.objects.get(dni=row[2].strip()),
                                                   'start_date': date,
                                                   'is_leadership': row[4] == 'SI'})
        except Exception:
            logger.warning("No se pudo importar el puesto %s", row[0], exc_info=True)


class ApprovalLevelCreate(CreateView):
    template_name = 'administration/approval_level_form.html'
    form_class = ApprovalLevelForm

    @method_decorator(
        requires('administration.add_approvallevel'))
    def dispatch(self, *args, **kwargs):
        return super(ApprovalLevelCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:approval_level_detail', args=[self.object.pk])


class ProfessionCreate(CreateView):
    template_name = 'administration/profession_form.html'
    form_class = ProfessionForm

    @method_decorator(requires('administration.add_profession'))
    def dispatch(self, *args, **kwargs):
        return super(ProfessionCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:profession_detail', args=[self.object.pk])


class OfficeCreate(CreateView):
    template_name = 'administration/office_form.html'
    form_class = OfficeForm

    @method_decorator(requires('administration.add_office'))
    def dispatch(self, *args, **kwargs):
        return super(OfficeCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:office_list')


class WorkerCreate(CreateView):
    template_name = 'administration/worker_form.html'
    form_class = WorkerForm

    @method_decorator(requires('administration.add_worker'))
    def dispatch(self, *args, **kwargs):
        return super(WorkerCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:worker_detail', args=[self.object.pk])


class ProducerCreate(CreateView):
    template_name = 'administration/producer_form.html'
    form_class = ProducerForm

    @method_decorator(requires('administration.add_producer'))
    def dispatch(self, *args, **kwargs):
        return super(ProducerCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:producer_detail', args=[self.object.pk])


class PositionCreate(CreateView):
    template_name = 'administration/position_form.html'
    form_class = PositionForm

    @method_decorator(requires('administration.add_position'))
    def dispatch(self, *args, **kwargs):
        return super(PositionCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:position_detail', args=[self.object.pk])


class OfficeDetail(DetailView):
    model = Office
    template_name = 'administration/office_detail.html'


class WorkerDetail(DetailView):
    model = Worker
    template_name = 'administration/worker_detail.html'


class ProducerDetail(DetailView):
    model = Producer
    template_name = 'administration/producer_detail.html'


class PositionDetail(DetailView):
    model = Position
    template_name = 'administration/position_detail.html'


class ProfessionDetail(DetailView):
    model = Profession
    template_name = 'administration/profession_detail.html'


class ApprovalLevelDetail(DetailView):
    model = ApprovalLevel
    template_name = 'administration/approval_level_detail.html'


class OfficeList(ListView):
    model = Office
    template_name = 'administration/office_list.html'
    context_object_name = 'offices'
    queryset = Office.objects.all().order_by('name')


class WorkerList(ListView):
    model = Worker
    template_name = 'administration/worker_list.html'
    context_object_name = 'workers'


class ProducerList(ListView):
    model = Producer
    template_name = 'administration/producer_list.html'
    context_object_name = 'productores'


class PositionList(ListView):
    model = Position
    template_name = 'administration/position_list.html'
    context_object_name = 'positions'
    queryset = Position.objects.filter(is_active=True)


class ProfessionList(ListView):
    model = Profession
    template_name = 'administration/profession_list.html'
    context_object_name = 'profesiones'


class ApprovalLevelList(ListView):
    model = ApprovalLevel
    template_name = 'administration/approval_level_list.html'
    context_object_name = 'niveles'


class ApprovalLevelUpdate(UpdateView):
    model = ApprovalLevel
    template_name = 'administration/approval_level_form.html'
    form_class = ApprovalLevelForm

    @method_decorator(
        requires('administration.change_approvallevel'))
    def dispatch(self, *args, **kwargs):
        return super(ApprovalLevelUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:approval_level_detail', args=[self.object.pk])


class ProfessionUpdate(UpdateView):
    model = Profession
    template_name = 'administration/profession_form.html'
    form_class = ProfessionForm

    @method_decorator(
        requires('administration.change_profession'))
    def dispatch(self, *args, **kwargs):
        return super(ProfessionUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:profession_detail', args=[self.object.pk])


class OfficeUpdate(UpdateView):
    model = Office
    template_name = 'administration/office_form.html'
    form_class = OfficeForm
    success_url = reverse_lazy('administration:office_list')

    @method_decorator(requires('administration.change_office'))
    def dispatch(self, *args, **kwargs):
        return super(OfficeUpdate, self).dispatch(*args, **kwargs)


class WorkerUpdate(UpdateView):
    model = Worker
    template_name = 'administration/worker_form.html'
    form_class = WorkerForm

    @method_decorator(
        requires('administration.change_worker'))
    def dispatch(self, *args, **kwargs):
        return super(WorkerUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:worker_detail', args=[self.object.pk])


class ProducerUpdate(UpdateView):
    model = Producer
    template_name = 'administration/producer_form.html'
    form_class = ProducerForm

    @method_decorator(
        requires('administration.change_producer'))
    def dispatch(self, *args, **kwargs):
        return super(ProducerUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('administration:producer_detail', args=[self.object.pk])


class PositionUpdate(UpdateView):
    model = Position
    template_name = 'administration/position_form.html'
    form_class = PositionUpdateForm

    @method_decorator(requires('administration.change_position'))
    def dispatch(self, *args, **kwargs):
        return super(PositionUpdate, self).dispatch(*args, **kwargs)

    def get_initial(self):
        initial = super(PositionUpdate, self).get_initial()
        initial['start_date'] = self.object.start_date.strftime('%d/%m/%Y')
        if self.object.end_date is not None:
            initial['end_date'] = self.object.end_date.strftime('%d/%m/%Y')
        return initial

    def get_success_url(self):
        return reverse('administration:position_detail', args=[self.object.pk])


class OfficeExcelReport(TemplateView):
    def get(self, request, *args, **kwargs):
        offices = Office.objects.filter(is_active=True).order_by('code')
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE OFICINAS'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'CODIGO'
        ws['C3'] = 'NOMBRE'
        ws['D3'] = 'DEPENDENCIA'
        ws['E3'] = 'GERENCIA'
        cont = 4
        for office in offices:
            try:
                ws.cell(row=cont, column=2).value = office.code
                ws.cell(row=cont, column=3).value = office.name
                ws.cell(row=cont, column=4).value = office.dependency.name
                ws.cell(row=cont, column=5).value = office.management.name
                cont = cont + 1
            except Exception:
                logger.warning("No se pudo exportar la oficina %s", office.pk, exc_info=True)
        file_name = "Oficinas.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class ProfessionExcelReport(TemplateView):
    def get(self, request, *args, **kwargs):
        profesiones = Profession.objects.filter(is_active=True)
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE PROFESIONES'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'ABREVIATURA'
        ws['C3'] = 'DESCRIPCION'
        ws['D3'] = 'ESTADO'
        cont = 4
        for profession in profesiones:
            ws.cell(row=cont, column=2).value = profession.abbreviation
            ws.cell(row=cont, column=3).value = profession.description
            ws.cell(row=cont, column=4).value = profession.is_active
            cont = cont + 1
        file_name = "Profesiones.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class PositionExcelReport(TemplateView):
    def get(self, request, *args, **kwargs):
        positions = Position.objects.filter(is_active=True)
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE PUESTOS'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'NOMBRE'
        ws['C3'] = 'OFICINA'
        ws['D3'] = 'TRABAJADOR'
        ws['E3'] = 'FECHA INICIO'
        ws['F3'] = 'FECHA FIN'
        ws['G3'] = 'ES JEFATURA'
        ws['H3'] = 'ESTADO'
        cont = 4
        for position in positions:
            ws.cell(row=cont, column=2).value = position.name
            ws.cell(row=cont, column=3).value = position.office.name
            ws.cell(row=cont, column=4).value = position.worker.full_name()
            ws.cell(row=cont, column=5).value = position.start_date.strftime('%d/%m/%Y')
            ws.cell(row=cont, column=6).value = position.end_date
            if position.is_leadership:
                ws.cell(row=cont, column=7).value = "SI"
            else:
                ws.cell(row=cont, column=7).value = "NO"
            ws.cell(row=cont, column=8).value = position.is_active
            cont = cont + 1
        file_name = "Puestos.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class WorkerExcelReport(TemplateView):
    def get(self, request, *args, **kwargs):
        workers = Worker.objects.filter(is_active=True)
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE TRABAJADORES'
        ws.merge_cells('B1:I1')
        ws['B3'] = 'USUARIO'
        ws['C3'] = 'DNI'
        ws['D3'] = 'APELLIDOS'
        ws['E3'] = 'NOMBRES'
        ws['F3'] = 'EMAIL'
        ws['G3'] = 'ESTADO'
        cont = 4
        for worker in workers:
            ws.cell(row=cont, column=2).value = worker.user.username
            ws.cell(row=cont, column=3).value = worker.dni
            ws.cell(row=cont, column=4).value = worker.last_name
            ws.cell(row=cont, column=5).value = worker.first_name
            ws.cell(row=cont, column=6).value = worker.user.email
            ws.cell(row=cont, column=7).value = worker.is_active
            cont = cont + 1
        file_name = "Trabajadores.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response
