# -*- coding: utf-8 -*- 
from django.utils import timezone
from django.views.generic.base import View, TemplateView
from django.views.generic.list import ListView

from purchases.models import Supplier, PurchaseOrder, PaymentMethod, PurchaseOrderDetail, RequirementDetail, ServiceOrder,\
    ServiceOrderDetail, ServiceConformity,\
    ServiceConformityDetail, QuotationDetail, Quotation
from django.views.generic.edit import FormView, UpdateView, CreateView
from purchases.forms import SupplierForm, QuotationForm, PurchaseOrderForm,\
    ServiceOrderForm, ServiceConformityForm, PurchaseOrderDetailFormSet,\
    ServiceOrderDetailFormSet, ServiceConformityDetailFormSet, QuotationDetailFormSet,\
    OrderDateReportForm
from django.urls import reverse_lazy, reverse
from django.http.response import HttpResponseRedirect
import json
from django.http import HttpResponse
import datetime
import simplejson
from openpyxl import Workbook
from django.views.generic.detail import DetailView
# from reportlab.lib.pagesizes import cm
import locale
from security.permissions import requires
from django.utils.decorators import method_decorator
from django.db import transaction, IntegrityError
from django.core.exceptions import ObjectDoesNotExist
from accounting.forms import UploadForm
from django.db.models import Q
from django.contrib import messages
from django.shortcuts import render, get_object_or_404

from products.models import Product, UnitOfMeasure, ProductGroup
from datetime import date
from purchases.reports import purchase_order_xls_report, PurchaseOrderPdf,\
    ServiceOrderPdf, ServiceConformityMemoPdf, QuotationRequestPdf
from tambox.config import configuration, purchase_tax
from tambox.views import CsvImportMixin, AjaxOnlyMixin, HtmxListMixin

locale.setlocale(locale.LC_ALL, "")


class Dashboard(View):

    def get(self, request, *args, **kwargs):
        notification_list = []
        supplier_count = Supplier.objects.count()
        product_count = Product.objects.filter(is_service=False).count()
        unit_of_measure_count = UnitOfMeasure.objects.count()
        supply_group_count = ProductGroup.objects.count()
        service_count = Product.objects.filter(is_service=True).count()
        unit_of_measure, creado = UnitOfMeasure.objects.get_or_create(code='SERV',
                                                                   defaults={'description': 'SERVICIO'})
        if supplier_count == 0:
            notification_list.append("No se ha creado ningún proveedor")
        if creado:
            notification_list.append("Se ha creado la unidad de medida SERVICIO")
        if product_count == 0:
            notification_list.append("No se ha creado ningún producto")
        if unit_of_measure_count == 0:
            notification_list.append("No se ha creado ningún tipo de unidad de medida")
        if supply_group_count == 0:
            notification_list.append("No se ha creado ningún grupo de productos")
        if service_count == 0:
            notification_list.append("No se ha creado ningún service")
        context = {'notifications': notification_list}
        return render(request, 'purchases/purchases_dashboard.html', context)


class QuotationSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('code',)

    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.GET['code']
            quotation = Quotation.objects.get(code=code)
            quotation_json = {}
            quotation_json['tax_id'] = quotation.supplier.tax_id
            quotation_json['business_name'] = quotation.supplier.business_name
            quotation_json['address'] = quotation.supplier.address
            data = simplejson.dumps(quotation_json)
            return HttpResponse(data, 'application/json')


class SupplierNameSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('business_name',)

    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            business_name = request.GET['business_name']
            suppliers = Supplier.objects.filter(business_name__icontains=business_name)[:20]
            supplier_list = []
            for supplier in suppliers:
                supplier_json = {}
                supplier_json['label'] = supplier.business_name
                supplier_json['tax_id'] = supplier.tax_id
                supplier_json['address'] = supplier.address
                supplier_json['is_service_provider'] = supplier.is_service_provider
                supplier_json['order'] = str(ServiceOrder.objects.last_record())
                supplier_list.append(supplier_json)
            data = json.dumps(supplier_list)
            return HttpResponse(data, 'application/json')


class SupplierTaxIdSearch(AjaxOnlyMixin, TemplateView):

    required_params = ('tax_id',)

    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            tax_id = request.GET['tax_id']
            supplier = Supplier.objects.get(tax_id=tax_id)
            supplier_json = {}
            supplier_json['business_name'] = supplier.business_name
            supplier_json['address'] = supplier.address
            supplier_json['status'] = supplier.sunat_status
            supplier_json['is_service_provider'] = supplier.is_service_provider
            supplier_json['order'] = str(ServiceOrder.objects.last_record())
            data = simplejson.dumps(supplier_json)
            return HttpResponse(data, 'application/json')


class SupplierImport(CsvImportMixin, FormView):
    template_name = 'purchases/supplier_upload.html'
    form_class = UploadForm
    success_url = reverse_lazy('purchases:supplier_list')

    def process_row(self, row):
        Supplier.objects.get_or_create(tax_id=row[0],
                                        defaults={'business_name': row[1],
                                                  'address': row[2],
                                                  'registration_date': datetime.datetime.now(),
                                                  'sunat_status': 'ACTIVO',
                                                  'sunat_condition': 'HABIDO',
                                                  'ciiu': 'CUALQUIERA'})


class SupplierCreate(CreateView):
    model = Supplier
    context_object_name = 'supplier'
    template_name = 'purchases/supplier_form.html'
    form_class = SupplierForm

    @method_decorator(requires('purchases.add_supplier'))
    def dispatch(self, *args, **kwargs):
        return super(SupplierCreate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('purchases:supplier_detail', args=[self.object.pk])

    def form_invalid(self, form):
        return self.render_to_response(self.get_context_data(form=form))


class QuotationCreate(CreateView):
    form_class = QuotationForm
    template_name = "purchases/quotation_form.html"
    model = Quotation
    context_object_name = 'quotation'

    @method_decorator(requires('purchases.add_quotation'))
    def dispatch(self, *args, **kwargs):
        return super(QuotationCreate, self).dispatch(*args, **kwargs)

    def get_initial(self):
        initial = super(QuotationCreate, self).get_initial()
        initial['date'] = date.today().strftime('%d/%m/%Y')
        return initial

    def get(self, request, *args, **kwargs):
        self.object = None
        suppliers = Supplier.objects.all()
        if not suppliers:
            return HttpResponseRedirect(reverse('purchases:supplier_create'))
        else:
            form_class = self.get_form_class()
            form = self.get_form(form_class)
            quotation_detail_formset = QuotationDetailFormSet()
            return self.render_to_response(self.get_context_data(form=form,
                                                                 quotation_detail_formset=quotation_detail_formset))

    def post(self, request, *args, **kwargs):
        self.object = None
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        quotation_detail_formset = QuotationDetailFormSet(request.POST)
        if form.is_valid() and quotation_detail_formset.is_valid():
            return self.form_valid(form, quotation_detail_formset)
        else:
            return self.form_invalid(form, quotation_detail_formset)

    def form_valid(self, form, quotation_detail_formset):
        # try:
        with transaction.atomic():
            self.object = form.save()
            reference = self.object.requirement
            details = []
            cont = 1
            for quotation_detail_form in quotation_detail_formset:
                requirement = quotation_detail_form.cleaned_data.get('requirement')
                quantity = quotation_detail_form.cleaned_data.get('quantity')
                requirement_detail = RequirementDetail.objects.get(pk=requirement)
                if quantity:
                    quotation_detail = QuotationDetail(requirement_detail=requirement_detail,
                                                           line_number=cont,
                                                           quotation=self.object,
                                                           quantity=quantity)
                    details.append(quotation_detail)

                    cont = cont + 1
            QuotationDetail.objects.bulk_create(details, reference, None)
            return HttpResponseRedirect(reverse('purchases:quotation_detail', args=[self.object.code]))
        # except IntegrityError:
        # messages.error(self.request, 'Error guardando la cotizacion.')

    def form_invalid(self, form, quotation_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             quotation_detail_formset=quotation_detail_formset))


class PurchaseOrderCreate(CreateView):
    form_class = PurchaseOrderForm
    template_name = "purchases/purchase_order_form.html"
    model = PurchaseOrder

    @method_decorator(requires('purchases.add_purchaseorder'))
    def dispatch(self, *args, **kwargs):
        return super(PurchaseOrderCreate, self).dispatch(*args, **kwargs)

    def get_initial(self):
        initial = super(PurchaseOrderCreate, self).get_initial()
        try:
            tax_amount = purchase_tax().amount
        except AttributeError:
            return HttpResponseRedirect(reverse('accounting:configuration'))
        initial['date'] = date.today().strftime('%d/%m/%Y')
        initial['code'] = PurchaseOrder.objects.last_record()
        initial['current_tax'] = tax_amount
        initial['total'] = 0
        initial['subtotal'] = 0
        initial['tax'] = 0
        initial['total_in_words'] = ''
        return initial

    def get(self, request, *args, **kwargs):
        self.object = None
        payment_methods = PaymentMethod.objects.all().order_by('description')
        if not payment_methods:
            return HttpResponseRedirect(reverse('accounting:payment_method_create'))
        else:
            try:
                configuration()
                form_class = self.get_form_class()
                form = self.get_form(form_class)
                purchase_order_detail_formset = PurchaseOrderDetailFormSet()
                return self.render_to_response(self.get_context_data(form=form,
                                                                     purchase_order_detail_formset=purchase_order_detail_formset))
            except Exception:
                return HttpResponseRedirect(reverse('accounting:configuration'))

    def post(self, request, *args, **kwargs):
        self.object = None
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        purchase_order_detail_formset = PurchaseOrderDetailFormSet(request.POST)
        if form.is_valid() and purchase_order_detail_formset.is_valid():
            return self.form_valid(form, purchase_order_detail_formset)
        else:
            return self.form_invalid(form, purchase_order_detail_formset)

    def form_valid(self, form, purchase_order_detail_formset):
        try:
            with transaction.atomic():
                self.object = form.save()
                reference = self.object.quotation
                details = []
                cont = 1
                for purchase_order_detail_form in purchase_order_detail_formset:
                    quotation = purchase_order_detail_form.cleaned_data.get('quotation')
                    code = purchase_order_detail_form.cleaned_data.get('code')
                    quantity = purchase_order_detail_form.cleaned_data.get('quantity')
                    price = purchase_order_detail_form.cleaned_data.get('price')
                    amount = purchase_order_detail_form.cleaned_data.get('amount')
                    tax = purchase_order_detail_form.cleaned_data.get('tax')
                    if quantity and price and amount and tax:
                        try:
                            quotation_detail = QuotationDetail.objects.get(pk=quotation)
                            purchase_order_detail = PurchaseOrderDetail(quotation_detail=quotation_detail,
                                                                      line_number=cont,
                                                                      order=self.object,
                                                                      quantity=quantity,
                                                                      price=price)
                        except QuotationDetail.DoesNotExist:
                            product = Product.objects.get(pk=code)
                            purchase_order_detail = PurchaseOrderDetail(product=product,
                                                                      line_number=cont,
                                                                      order=self.object,
                                                                      quantity=quantity,
                                                                      price=price)
                        details.append(purchase_order_detail)
                        cont = cont + 1
                if cont > 1:
                    PurchaseOrderDetail.objects.bulk_create(details, reference)
                return HttpResponseRedirect(reverse('purchases:purchase_order_detail', args=[self.object.pk]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando la orden de compra.')

    def form_invalid(self, form, purchase_order_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             purchase_order_detail_formset=purchase_order_detail_formset))


class ServiceOrderCreate(CreateView):
    form_class = ServiceOrderForm
    template_name = "purchases/service_order_form.html"
    model = ServiceOrder

    @method_decorator(requires('purchases.add_serviceorder'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceOrderCreate, self).dispatch(*args, **kwargs)

    def get_initial(self):
        initial = super(ServiceOrderCreate, self).get_initial()
        initial['code'] = ServiceOrder.objects.last_record()
        initial['total'] = 0
        initial['subtotal'] = 0
        initial['tax'] = 0
        initial['total_in_words'] = ''
        return initial

    def get(self, request, *args, **kwargs):
        self.object = None
        payment_methods = PaymentMethod.objects.all().order_by('description')
        if not payment_methods:
            return HttpResponseRedirect(reverse('accounting:payment_method_create'))
        else:
            form_class = self.get_form_class()
            form = self.get_form(form_class)
            service_order_detail_formset = ServiceOrderDetailFormSet()
            return self.render_to_response(self.get_context_data(form=form,
                                                                 service_order_detail_formset=service_order_detail_formset))

    def post(self, request, *args, **kwargs):
        self.object = None
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        service_order_detail_formset = ServiceOrderDetailFormSet(request.POST)
        if form.is_valid() and service_order_detail_formset.is_valid():
            return self.form_valid(form, service_order_detail_formset)
        else:
            return self.form_invalid(form, service_order_detail_formset)

    def form_valid(self, form, service_order_detail_formset):
        try:
            with transaction.atomic():
                self.object = form.save()
                reference = self.object.quotation
                details = []
                cont = 1
                for service_order_detail_form in service_order_detail_formset:
                    quotation = service_order_detail_form.cleaned_data.get('quotation')
                    code = service_order_detail_form.cleaned_data.get('code')
                    quantity = service_order_detail_form.cleaned_data.get('quantity')
                    price = service_order_detail_form.cleaned_data.get('price')
                    amount = service_order_detail_form.cleaned_data.get('amount')
                    if quantity and price and amount:
                        try:
                            quotation_detail = QuotationDetail.objects.get(pk=quotation)
                            service_order_detail = ServiceOrderDetail(quotation_detail=quotation_detail,
                                                                            line_number=cont,
                                                                            order=self.object,
                                                                            quantity=quantity,
                                                                            price=price)
                        except QuotationDetail.DoesNotExist:
                            product = Product.objects.get(pk=code)
                            service_order_detail = ServiceOrderDetail(product=product,
                                                                            line_number=cont,
                                                                            order=self.object,
                                                                            quantity=quantity,
                                                                            price=price)

                        details.append(service_order_detail)
                        cont = cont + 1
                ServiceOrderDetail.objects.bulk_create(details, reference)
                return HttpResponseRedirect(reverse('purchases:service_order_detail', args=[self.object.code]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando la cotizacion.')

    def form_invalid(self, form, service_order_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             service_order_detail_formset=service_order_detail_formset))


class ServiceConformityCreate(CreateView):
    form_class = ServiceConformityForm
    template_name = "purchases/service_conformity_form.html"
    model = ServiceConformity

    @method_decorator(requires('purchases.add_serviceconformity'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceConformityCreate, self).dispatch(*args, **kwargs)

    def get_initial(self):
        initial = super(ServiceConformityCreate, self).get_initial()
        initial['total'] = 0
        initial['subtotal'] = 0
        initial['total_in_words'] = ''
        return initial

    def get(self, request, *args, **kwargs):
        self.object = None
        payment_methods = PaymentMethod.objects.all().order_by('description')
        if not payment_methods:
            return HttpResponseRedirect(reverse('accounting:payment_method_create'))
        else:
            form_class = self.get_form_class()
            form = self.get_form(form_class)
            service_conformity_detail_formset = ServiceConformityDetailFormSet()
            return self.render_to_response(self.get_context_data(form=form,
                                                                 service_conformity_detail_formset=service_conformity_detail_formset))

    def post(self, request, *args, **kwargs):
        self.object = None
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        service_conformity_detail_formset = ServiceConformityDetailFormSet(request.POST)
        if form.is_valid() and service_conformity_detail_formset.is_valid():
            return self.form_valid(form, service_conformity_detail_formset)
        else:
            return self.form_invalid(form, service_conformity_detail_formset)

    def form_valid(self, form, service_conformity_detail_formset):
        try:
            with transaction.atomic():
                self.object = form.save()
                reference = self.object.service_order
                details = []
                cont = 1
                for service_order_detail_form in service_conformity_detail_formset:
                    service_order = service_order_detail_form.cleaned_data.get('service_order')
                    quantity = service_order_detail_form.cleaned_data.get('quantity')
                    price = service_order_detail_form.cleaned_data.get('price')
                    amount = service_order_detail_form.cleaned_data.get('amount')
                    service_order_detail = ServiceOrderDetail.objects.get(pk=service_order)
                    if quantity and price and amount:
                        service_conformity_detail = ServiceConformityDetail(
                            service_order_detail=service_order_detail,
                            line_number=cont,
                            conformity=self.object,
                            quantity=quantity)
                        details.append(service_conformity_detail)
                        cont = cont + 1
                ServiceConformityDetail.objects.bulk_create(details, reference)
                return HttpResponseRedirect(reverse('purchases:service_conformity_detail_view', args=[self.object.code]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando la cotizacion.')

    def form_invalid(self, form, service_conformity_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             service_conformity_detail_form=service_conformity_detail_formset))


class SupplierDetail(DetailView):
    model = Supplier
    template_name = 'purchases/supplier_detail.html'


class QuotationDetailView(DetailView):
    model = Quotation
    context_object_name = 'quotation'
    template_name = 'purchases/quotation_detail.html'


class PurchaseOrderDetailView(DetailView):
    model = PurchaseOrder
    template_name = 'purchases/purchase_order_detail.html'


class ServiceOrderDetailView(DetailView):
    model = ServiceOrder
    template_name = 'purchases/service_order_detail.html'


class ServiceConformityDetailView(DetailView):
    model = ServiceConformity
    template_name = 'purchases/service_conformity_detail.html'


class QuotationDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('purchases.delete_quotation'))
    def dispatch(self, *args, **kwargs):
        return super(QuotationDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            quotation = Quotation.objects.get(code=code)
            quotation_json = {}
            quotation_json['code'] = code
            service_orders = quotation.service_orders.all()
            if len(service_orders) > 0:
                quotation_json['orders'] = 'SI'
            else:
                quotation_json['orders'] = 'NO'
                purchase_orders = quotation.purchase_orders.all()
                if len(purchase_orders) > 0:
                    quotation_json['orders'] = 'SI'
                else:
                    quotation_json['orders'] = 'NO'

                with transaction.atomic():
                    quotation.delete_reference()
                    quotation.delete_quotation()
                    QuotationDetail.objects.filter(quotation=quotation).delete()
            data = simplejson.dumps(quotation_json)
            return HttpResponse(data, 'application/json')


class PurchaseOrderDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('purchases.delete_purchaseorder'))
    def dispatch(self, *args, **kwargs):
        return super(PurchaseOrderDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            order = PurchaseOrder.objects.get(code=code)
            movement_json = {}
            movement_json['code'] = code
            if len(order.movements.all()) > 0:
                movement_json['movements'] = 'SI'
            else:
                movement_json['movements'] = 'NO'
                with transaction.atomic():
                    if order.quotation is not None:
                        order.delete_reference()
                    PurchaseOrder.objects.filter(code=code).update(status=PurchaseOrder.STATUS.CANC, quotation=None)
                    PurchaseOrderDetail.objects.filter(order=order).delete()
            data = simplejson.dumps(movement_json)
            return HttpResponse(data, 'application/json')


class ServiceOrderDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('purchases.delete_serviceorder'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceOrderDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            order = ServiceOrder.objects.get(code=code)
            order_json = {}
            order_json['code'] = code
            if len(order.conformities.all()) > 0:
                order_json['conformidades'] = 'SI'
            else:
                order_json['conformidades'] = 'NO'
                with transaction.atomic():
                    if order.quotation is not None:
                        order.delete_reference()
                    ServiceOrder.objects.filter(code=code).update(status=ServiceOrder.STATUS.CANC,
                                                                        quotation=None)
                    ServiceOrderDetail.objects.filter(order=order).delete()
            data = simplejson.dumps(order_json)
            return HttpResponse(data, 'application/json')


class ServiceConformityDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('purchases.delete_serviceconformity'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceConformityDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            conformity = ServiceConformity.objects.get(code=code)
            conformity_json = {}
            conformity_json['code'] = code
            with transaction.atomic():
                if conformity.service_order is not None:
                    conformity.delete_reference()
                ServiceConformity.objects.filter(code=code).update(is_active=False)
                ServiceConformityDetail.objects.filter(conformity=conformity).delete()
            data = simplejson.dumps(conformity_json)
            return HttpResponse(data, 'application/json')


class SupplierDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('purchases.delete_supplier'))
    def dispatch(self, *args, **kwargs):
        return super(SupplierDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            tax_id = request.POST['tax_id']
            supplier_json = {}
            supplier_json['tax_id'] = tax_id
            Supplier.objects.filter(tax_id=tax_id).update(is_active=False)
            data = simplejson.dumps(supplier_json)
            return HttpResponse(data, 'application/json')


class SupplierList(HtmxListMixin, ListView):
    model = Supplier
    template_name = 'purchases/supplier_list.html'
    fragment_template_name = 'purchases/includes/supplier_rows.html'
    context_object_name = 'suppliers'
    paginate_by = 10
    queryset = Supplier.objects.filter(is_active=True).order_by('business_name')
    search_fields = ('tax_id', 'business_name')

    @method_decorator(requires('purchases.ver_tabla_proveedores'))
    def dispatch(self, *args, **kwargs):
        return super(SupplierList, self).dispatch(*args, **kwargs)


class QuotationList(HtmxListMixin, ListView):
    model = Quotation
    template_name = 'purchases/quotation_list.html'
    fragment_template_name = 'purchases/includes/quotation_rows.html'
    context_object_name = 'quotations'
    paginate_by = 10
    queryset = Quotation.objects.exclude(status=Quotation.STATUS.CANC).order_by('code')
    search_fields = ('code', 'supplier__business_name', 'requirement__code')

    @method_decorator(requires('purchases.ver_tabla_cotizaciones'))
    def dispatch(self, *args, **kwargs):
        return super(QuotationList, self).dispatch(*args, **kwargs)


class PurchaseOrderList(HtmxListMixin, ListView):
    model = PurchaseOrder
    template_name = 'purchases/purchase_order_list.html'
    fragment_template_name = 'purchases/includes/purchase_order_rows.html'
    context_object_name = 'purchase_orders'
    paginate_by = 10
    queryset = PurchaseOrder.objects.exclude(status=PurchaseOrder.STATUS.CANC).order_by('code')
    search_fields = ('code', 'quotation__code')

    @method_decorator(
        requires('purchases.ver_tabla_ordenes_compra'))
    def dispatch(self, *args, **kwargs):
        return super(PurchaseOrderList, self).dispatch(*args, **kwargs)


class ServiceOrderList(HtmxListMixin, ListView):
    model = ServiceOrder
    template_name = 'purchases/service_order_list.html'
    fragment_template_name = 'purchases/includes/service_order_rows.html'
    context_object_name = 'service_orders'
    paginate_by = 10
    queryset = ServiceOrder.objects.filter().order_by('code')
    search_fields = ('code', 'quotation__code', 'supplier__business_name')

    @method_decorator(
        requires('purchases.ver_tabla_ordenes_servicios'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceOrderList, self).dispatch(*args, **kwargs)


class PurchaseOrderListByQuotation(HtmxListMixin, ListView):
    model = PurchaseOrder
    template_name = 'purchases/purchase_order_list.html'
    fragment_template_name = 'purchases/includes/purchase_order_rows.html'
    context_object_name = 'purchase_orders'
    paginate_by = 10
    search_fields = ('code',)

    @method_decorator(
        requires('purchases.ver_tabla_ordenes_compra'))
    def dispatch(self, *args, **kwargs):
        return super(PurchaseOrderListByQuotation, self).dispatch(*args, **kwargs)

    def get_queryset(self):
        quotation = Quotation.objects.get(pk=self.kwargs['quotation'])
        queryset = quotation.purchase_orders.all()
        return queryset


class ServiceOrderListByQuotation(HtmxListMixin, ListView):
    model = PurchaseOrder
    template_name = 'purchases/service_order_list.html'
    fragment_template_name = 'purchases/includes/service_order_rows.html'
    context_object_name = 'service_orders'
    paginate_by = 10
    search_fields = ('code',)

    @method_decorator(
        requires('purchases.ver_tabla_ordenes_servicios'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceOrderListByQuotation, self).dispatch(*args, **kwargs)

    def get_queryset(self):
        quotation = Quotation.objects.get(pk=self.kwargs['quotation'])
        queryset = quotation.service_orders.all()
        return queryset


class ServiceConformityList(HtmxListMixin, ListView):
    model = ServiceConformity
    template_name = 'purchases/service_conformity_list.html'
    fragment_template_name = 'purchases/includes/service_conformity_rows.html'
    context_object_name = 'conformities'
    paginate_by = 10
    queryset = ServiceConformity.objects.filter(is_active=True).order_by('code')
    search_fields = ('code', 'service_order__code')

    @method_decorator(
        requires('purchases.ver_tabla_conformidades_servicio'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceConformityList, self).dispatch(*args, **kwargs)


class MovementListByPurchaseOrder(ListView):
    template_name = 'warehouse/movements.html'
    context_object_name = 'movements'

    @method_decorator(requires('warehouse.ver_tabla_movimientos'))
    def dispatch(self, *args, **kwargs):
        return super(MovementListByPurchaseOrder, self).dispatch(*args, **kwargs)

    def get_queryset(self):
        purchase_order = PurchaseOrder.objects.get(pk=self.kwargs['order'])
        queryset = purchase_order.movements.all()
        return queryset


class ServiceConformityListByServiceOrder(ListView):
    template_name = 'purchases/service_conformity_list.html'
    context_object_name = 'conformities'

    @method_decorator(
        requires('purchases.ver_tabla_conformidades_servicio'))
    def dispatch(self, *args, **kwargs):
        return super(ServiceConformityListByServiceOrder, self).dispatch(*args, **kwargs)

    def get_queryset(self):
        service_order = ServiceOrder.objects.get(pk=self.kwargs['order'])
        queryset = service_order.conformities.all()
        return queryset


class SupplierUpdate(UpdateView):
    model = Supplier
    context_object_name = 'supplier'
    template_name = 'purchases/supplier_form.html'
    form_class = SupplierForm

    @method_decorator(requires('purchases.change_supplier'))
    def dispatch(self, *args, **kwargs):
        return super(SupplierUpdate, self).dispatch(*args, **kwargs)

    def get_success_url(self):
        return reverse('purchases:supplier_detail', args=[self.object.pk])

    def get_initial(self):
        initial = super(SupplierUpdate, self).get_initial()
        initial['registration_date'] = self.object.registration_date.strftime('%d/%m/%Y')
        return initial


class QuotationUpdate(UpdateView):
    form_class = QuotationForm
    template_name = "purchases/quotation_form.html"
    model = Quotation
    context_object_name = 'quotation'

    @method_decorator(requires('purchases.change_quotation'))
    def dispatch(self, *args, **kwargs):
        quotation = self.get_object()
        if quotation.status == Quotation.STATUS.PEND:
            return super(QuotationUpdate, self).dispatch(*args, **kwargs)
        else:
            return HttpResponseRedirect(reverse('security:permission_denied'))

    def get_initial(self):
        initial = super(QuotationUpdate, self).get_initial()
        quotation = self.object
        initial['code'] = quotation.code
        initial['tax_id'] = quotation.supplier.tax_id
        initial['business_name'] = quotation.supplier.business_name
        initial['address'] = quotation.supplier.address
        initial['date'] = quotation.date.strftime('%d/%m/%Y')
        initial['reference'] = quotation.requirement
        initial['notes'] = quotation.notes
        return initial

    def get_context_data(self, **kwargs):
        quotation = self.object
        details = QuotationDetail.objects.filter(quotation=quotation).order_by('line_number')
        detail_count = details.count()
        context = super(QuotationUpdate, self).get_context_data(**kwargs)
        context['quotation'] = quotation
        context['details'] = details
        context['detail_count'] = detail_count
        return context

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        details = QuotationDetail.objects.filter(quotation=self.object).order_by('line_number')
        details_data = []
        for detail in details:
            d = {'requirement': detail.requirement_detail.pk,
                 'code': detail.requirement_detail.product.code,
                 'name': detail.requirement_detail.product.description,
                 'unit': detail.requirement_detail.product.unit_of_measure.code,
                 'quantity': detail.quantity}
            details_data.append(d)
        quotation_detail_formset = QuotationDetailFormSet(initial=details_data)
        return self.render_to_response(self.get_context_data(form=form,
                                                             quotation_detail_formset=quotation_detail_formset))

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        quotation_detail_formset = QuotationDetailFormSet(request.POST)
        if form.is_valid() and quotation_detail_formset.is_valid():
            return self.form_valid(form, quotation_detail_formset)
        else:
            return self.form_invalid(form, quotation_detail_formset)

    def form_valid(self, form, quotation_detail_formset):
        try:
            with transaction.atomic():
                self.object.delete_reference()
                form.save()
                details = []
                cont = 1
                for quotation_detail_form in quotation_detail_formset:
                    requirement_detail = quotation_detail_form.cleaned_data.get('requirement')
                    quantity = quotation_detail_form.cleaned_data.get('quantity')
                    requirement_detail = RequirementDetail.objects.get(pk=requirement_detail)
                    if quantity:
                        quotation_detail = QuotationDetail(requirement_detail=requirement_detail,
                                                               line_number=cont,
                                                               quotation=self.object,
                                                               quantity=quantity)
                        details.append(quotation_detail)
                        cont = cont + 1
                QuotationDetail.objects.bulk_create(details, self.object.requirement, None)
                return HttpResponseRedirect(reverse('purchases:quotation_detail', args=[self.object.code]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando el requerimiento.')

    def form_invalid(self, form, quotation_detail_formset):
        return self.render_to_response(self.get_context_data(form=form))


class ServiceConformityUpdate(UpdateView):
    template_name = 'purchases/service_conformity_form.html'
    form_class = ServiceConformityForm
    model = ServiceConformity

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        return self.render_to_response(self.get_context_data(form=form))

    def get_initial(self):
        initial = super(ServiceConformityUpdate, self).get_initial()
        conformity = self.object
        initial['cod_conformidad_servicio'] = conformity.code
        initial['service_order'] = conformity.service_order
        initial['supporting_document'] = conformity.supporting_document
        initial['date'] = conformity.date.strftime('%d/%m/%Y')
        return initial

    def get_context_data(self, **kwargs):
        conformity = self.object
        details = ServiceConformityDetail.objects.filter(conformity=conformity)
        detail_count = details.count()
        context = super(ServiceConformityUpdate, self).get_context_data(**kwargs)
        context['conformity'] = conformity
        context['details'] = details
        context['detail_count'] = detail_count
        return context

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        if form.is_valid():
            return self.form_valid(form)
        else:
            return self.form_invalid(form)

    def form_valid(self, form):
        try:
            with transaction.atomic():
                self.object.delete_reference()
                form.save()
                return HttpResponseRedirect(reverse('purchases:quotation_detail', args=[self.object.code]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando el requerimiento.')


class PurchaseOrderUpdate(UpdateView):
    template_name = 'purchases/purchase_order_form.html'
    form_class = PurchaseOrderForm
    model = PurchaseOrder

    @method_decorator(requires('purchases.change_purchaseorder'))
    def dispatch(self, *args, **kwargs):
        purchase_order = self.get_object()
        if purchase_order.status == PurchaseOrder.STATUS.PEND:
            return super(PurchaseOrderUpdate, self).dispatch(*args, **kwargs)
        else:
            return HttpResponseRedirect(reverse('security:permission_denied'))

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        if self.object.status == PurchaseOrder.STATUS.PEND:
            form_class = self.get_form_class()
            form = self.get_form(form_class)
            details = (PurchaseOrderDetail.objects.filter(order=self.object)
                        .select_related('order', 'product__unit_of_measure',
                                        'quotation_detail__requirement_detail__product__unit_of_measure')
                        .order_by('line_number'))
            details_data = []
            for detail in details:
                try:
                    d = {'quotation': detail.quotation_detail.pk,
                         'code': detail.quotation_detail.requirement_detail.product.code,
                         'name': detail.quotation_detail.requirement_detail.product.description,
                         'unit': detail.quotation_detail.requirement_detail.product.unit_of_measure.code,
                         'quantity': detail.quantity,
                         'price': detail.price,
                         'tax': detail.tax,
                         'amount': detail.amount}
                except (ObjectDoesNotExist, AttributeError):
                    d = {'quotation': '0',
                         'code': detail.product.code,
                         'name': detail.product.description,
                         'unit': detail.product.unit_of_measure.code,
                         'quantity': detail.quantity,
                         'price': detail.price,
                         'tax': detail.tax,
                         'amount': detail.amount_without_tax}
                details_data.append(d)
            purchase_order_detail_formset = PurchaseOrderDetailFormSet(initial=details_data)
            return self.render_to_response(self.get_context_data(form=form,
                                                                 purchase_order_detail_formset=purchase_order_detail_formset))
        else:
            return HttpResponseRedirect(reverse('purchases:purchase_order_list'))

    def get_initial(self):
        initial = super(PurchaseOrderUpdate, self).get_initial()
        order = self.object
        quotation = order.quotation
        if quotation is None:
            supplier = order.supplier
        else:
            supplier = order.quotation.supplier
        initial['code'] = order.code
        initial['tax_id'] = supplier.tax_id
        initial['business_name'] = supplier.business_name
        initial['address'] = supplier.address
        initial['date'] = order.date.strftime('%d/%m/%Y')
        initial['payment_methods'] = order.payment_method
        initial['reference'] = order.quotation
        try:
            tax_amount = purchase_tax().amount
        except AttributeError:
            return HttpResponseRedirect(reverse('accounting:configuration'))
        initial['current_tax'] = tax_amount
        initial['total'] = order.total
        initial['subtotal'] = order.subtotal
        initial['tax'] = order.tax
        initial['total_in_words'] = order.total_in_words
        initial['notes'] = order.notes
        return initial

    def get_context_data(self, **kwargs):
        order = self.object
        context = super(PurchaseOrderUpdate, self).get_context_data(**kwargs)
        context['order'] = order
        return context

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        purchase_order_detail_formset = PurchaseOrderDetailFormSet(request.POST)
        if form.is_valid() and purchase_order_detail_formset.is_valid():
            return self.form_valid(form, purchase_order_detail_formset)
        else:
            return self.form_invalid(form, purchase_order_detail_formset)

    def form_valid(self, form, purchase_order_detail_formset):
        try:
            with transaction.atomic():
                if self.object.quotation is not None:
                    self.object.delete_reference()
                PurchaseOrderDetail.objects.filter(order=self.object).delete()
                self.object = form.save()
                reference = self.object.quotation
                details = []
                cont = 1
                for purchase_order_detail_form in purchase_order_detail_formset:
                    quotation = purchase_order_detail_form.cleaned_data.get('quotation')
                    code = purchase_order_detail_form.cleaned_data.get('code')
                    quantity = purchase_order_detail_form.cleaned_data.get('quantity')
                    price = purchase_order_detail_form.cleaned_data.get('price')
                    amount = purchase_order_detail_form.cleaned_data.get('amount')
                    tax = purchase_order_detail_form.cleaned_data.get('tax')
                    if quantity and price and amount and tax:
                        try:
                            quotation_detail = QuotationDetail.objects.get(pk=quotation)
                            purchase_order_detail = PurchaseOrderDetail(quotation_detail=quotation_detail,
                                                                      line_number=cont,
                                                                      order=self.object,
                                                                      quantity=quantity,
                                                                      price=price)
                        except QuotationDetail.DoesNotExist:
                            product = Product.objects.get(pk=code)
                            purchase_order_detail = PurchaseOrderDetail(product=product,
                                                                      line_number=cont,
                                                                      order=self.object,
                                                                      quantity=quantity,
                                                                      price=price)
                        details.append(purchase_order_detail)
                        cont = cont + 1
                        if cont > 1:
                            PurchaseOrderDetail.objects.bulk_create(details, reference)
                return HttpResponseRedirect(reverse('purchases:purchase_order_detail', args=[self.object.pk]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando la cotizacion.')

    def form_invalid(self, form, purchase_order_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             purchase_order_detail_formset=purchase_order_detail_formset))


class ServiceOrderUpdate(UpdateView):
    template_name = 'purchases/service_order_form.html'
    form_class = ServiceOrderForm
    model = ServiceOrder

    @method_decorator(requires('purchases.change_serviceorder'))
    def dispatch(self, *args, **kwargs):
        service_order = self.get_object()
        if service_order.status == ServiceOrder.STATUS.PEND:
            return super(ServiceOrderUpdate, self).dispatch(*args, **kwargs)
        else:
            return HttpResponseRedirect(reverse('security:permission_denied'))

    def get_initial(self):
        initial = super(ServiceOrderUpdate, self).get_initial()
        order = self.object
        quotation = order.quotation
        if quotation is None:
            supplier = order.supplier
        else:
            supplier = order.quotation.supplier
        initial['code'] = order.code
        initial['tax_id'] = supplier.tax_id
        initial['business_name'] = supplier.business_name
        initial['address'] = supplier.address
        initial['date'] = order.date.strftime('%d/%m/%Y')
        initial['payment_methods'] = order.payment_method
        initial['reference'] = order.quotation
        initial['process'] = order.process
        initial['total'] = order.total
        initial['subtotal'] = order.subtotal
        initial['tax'] = order.tax
        initial['total_in_words'] = order.total_in_words
        initial['notes'] = order.notes
        return initial

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        if self.object.status == PurchaseOrder.STATUS.PEND:
            form_class = self.get_form_class()
            form = self.get_form(form_class)
            details = ServiceOrderDetail.objects.filter(order=self.object).order_by('line_number')
            details_data = []
            for detail in details:
                try:
                    d = {'quotation': detail.quotation_detail.pk,
                         'code': detail.quotation_detail.requirement_detail.product.code,
                         'name': detail.quotation_detail.requirement_detail.product.description,
                         'unit': detail.quotation_detail.requirement_detail.product.unit_of_measure.code,
                         'quantity': detail.quantity,
                         'price': detail.price,
                         'amount': detail.amount}
                except (ObjectDoesNotExist, AttributeError):
                    d = {'quotation': '0',
                         'code': detail.product.code,
                         'name': detail.product.description,
                         'unit': detail.product.unit_of_measure.code,
                         'quantity': detail.quantity,
                         'price': detail.price,
                         'amount': detail.amount}
                details_data.append(d)
            service_order_detail_formset = ServiceOrderDetailFormSet(initial=details_data)
            return self.render_to_response(self.get_context_data(form=form,
                                                                 service_order_detail_formset=service_order_detail_formset))
        else:
            return HttpResponseRedirect(reverse('purchases:purchase_order_list'))

    def get_context_data(self, **kwargs):
        order = self.object
        context = super(ServiceOrderUpdate, self).get_context_data(**kwargs)
        context['order'] = order
        return context

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        service_order_detail_formset = ServiceOrderDetailFormSet(request.POST)
        if form.is_valid() and service_order_detail_formset.is_valid():
            return self.form_valid(form, service_order_detail_formset)
        else:
            return self.form_invalid(form, service_order_detail_formset)

    def form_valid(self, form, service_order_detail_formset):
        # try:
        with transaction.atomic():
            if self.object.quotation is not None:
                self.object.delete_reference()
            ServiceOrderDetail.objects.filter(order=self.object).delete()
            self.object = form.save()
            reference = self.object.quotation
            details = []
            cont = 1
            for service_order_detail_form in service_order_detail_formset:
                quotation = service_order_detail_form.cleaned_data.get('quotation')
                code = service_order_detail_form.cleaned_data.get('code')
                quantity = service_order_detail_form.cleaned_data.get('quantity')
                price = service_order_detail_form.cleaned_data.get('price')
                amount = service_order_detail_form.cleaned_data.get('amount')
                if quantity and price and amount:
                    try:
                        quotation_detail = QuotationDetail.objects.get(pk=quotation)
                        service_order_detail = ServiceOrderDetail(quotation_detail=quotation_detail,
                                                                        line_number=cont,
                                                                        order=self.object,
                                                                        quantity=quantity,
                                                                        price=price)
                    except ObjectDoesNotExist:
                        product = Product.objects.get(pk=code)
                        service_order_detail = ServiceOrderDetail(product=product,
                                                                        line_number=cont,
                                                                        order=self.object,
                                                                        quantity=quantity,
                                                                        price=price)
                    details.append(service_order_detail)
                    cont = cont + 1
            ServiceOrderDetail.objects.bulk_create(details, reference)
            return HttpResponseRedirect(reverse('purchases:service_order_detail', args=[self.object.code]))
            # except IntegrityError:
            # messages.error(self.request, 'Error guardando la Orden de Servicios.')

    def form_invalid(self, form, service_order_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             service_order_detail_formset=service_order_detail_formset))


class QuotationDetailRows(TemplateView):
    """Filas del formset de cotizacion para un requerimiento, como HTML.

    htmx las inserta en la tabla del formulario. Al venir del servidor traen los
    inputs reales del formset (incluido el `requirement` oculto), que es justo lo
    que la version anterior perdia al construir las filas a mano: el POST llegaba
    con `TOTAL_FORMS` pero sin ningun dato y el formset no validaba.
    """

    template_name = 'purchases/includes/quotation_detail_formset.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        code = self.request.GET.get('requirement', '')
        details = RequirementDetail.objects.filter(
            Q(status=RequirementDetail.STATUS.PEND) | Q(status=RequirementDetail.STATUS.COTIZ),
            requirement__code=code, product__isnull=False).order_by('line_number')
        details_data = []
        for detail in details:
            details_data.append({
                'requirement': detail.pk,
                'code': detail.product.code,
                'name': detail.product.description,
                'unit': detail.product.unit_of_measure.code,
                'quantity': detail.quantity - detail.served_quantity,
            })
        context['quotation_detail_formset'] = QuotationDetailFormSet(initial=details_data)
        return context


class PurchaseOrderDetailRows(TemplateView):
    """Filas del formset de orden de compra para una cotizacion, como HTML.

    htmx las inserta en la tabla; vienen del servidor con los inputs reales del
    formset, en lugar de armarse en el navegador desde el JSON.
    """

    template_name = 'purchases/includes/purchase_order_detail_formset.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        quotation_code = self.request.GET.get('quotation', '')
        details = QuotationDetail.objects.filter(
            Q(status=QuotationDetail.STATUS.PEND) | Q(status=QuotationDetail.STATUS.ELEG_PARC),
            quotation__code=quotation_code,
            requirement_detail__product__is_service=False).order_by('line_number')
        try:
            tax_amount = purchase_tax().amount
        except AttributeError:
            tax_amount = 0
        initial = []
        for detail in details:
            product = detail.requirement_detail.product
            quantity = detail.quantity - detail.requirement_detail.purchased_quantity
            amount = product.price * quantity
            base = amount / (tax_amount + 1)
            initial.append({'quotation': detail.pk,
                            'code': product.code,
                            'name': product.description,
                            'unit': product.unit_of_measure.code,
                            'quantity': quantity,
                            'price': product.price,
                            'tax': round(amount - base, 5),
                            'amount': round(amount, 5)})
        context['purchase_order_detail_formset'] = PurchaseOrderDetailFormSet(initial=initial)
        return context


class PurchaseOrderDetailRow(TemplateView):
    """Una fila vacia del formset de orden de compra, en el indice pedido.

    El navegador la agrega al final de la tabla y sube `TOTAL_FORMS`.
    """

    template_name = 'purchases/includes/purchase_order_detail_row.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        index = self.request.GET.get('index', '0')
        formset = PurchaseOrderDetailFormSet(
            initial=[{'quotation': '0', 'code': '', 'name': '', 'unit': '',
                      'quantity': 0, 'price': 0, 'tax': 0, 'amount': 0}])
        form = formset.forms[0]
        form.prefix = 'form-%s' % index
        context['form'] = form
        context['index'] = index
        return context


class ServiceOrderDetailRows(TemplateView):
    """Filas del formset de orden de servicio para una cotizacion (servicios)."""

    template_name = 'purchases/includes/service_order_detail_formset.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        quotation_code = self.request.GET.get('quotation', '')
        details = QuotationDetail.objects.filter(
            quotation__code=quotation_code,
            requirement_detail__product__is_service=True).order_by('line_number')
        initial = []
        for detail in details:
            product = detail.requirement_detail.product
            quantity = detail.quantity - detail.requirement_detail.purchased_quantity
            amount = product.price * quantity
            initial.append({'quotation': detail.pk,
                            'code': product.code,
                            'name': product.description,
                            'unit': product.unit_of_measure.code,
                            'quantity': quantity,
                            'price': product.price,
                            'amount': round(amount)})
        context['service_order_detail_formset'] = ServiceOrderDetailFormSet(initial=initial)
        return context


class ServiceOrderDetailRow(TemplateView):
    """Una fila vacia del formset de orden de servicio, en el indice pedido."""

    template_name = 'purchases/includes/service_order_detail_row.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        index = self.request.GET.get('index', '0')
        formset = ServiceOrderDetailFormSet(
            initial=[{'quotation': '0', 'code': '', 'name': '', 'unit': '',
                      'quantity': 0, 'price': 0, 'amount': 0}])
        form = formset.forms[0]
        form.prefix = 'form-%s' % index
        context['form'] = form
        context['index'] = index
        return context


class ServiceConformityDetailRows(TemplateView):
    """Filas del formset de conformidad de servicio para una orden de servicio."""

    template_name = 'purchases/includes/service_conformity_detail_formset.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        code = self.request.GET.get('service_order', '')
        details = ServiceOrderDetail.objects.filter(
            order__code=code,
            status=ServiceOrderDetail.STATUS.PEND).order_by('line_number')
        initial = []
        for detail in details:
            try:
                service = detail.quotation_detail.requirement_detail.product.description
                use = detail.quotation_detail.requirement_detail.use
            except (ObjectDoesNotExist, AttributeError):
                service = detail.product.description
                use = detail.product.unit_of_measure.description
            initial.append({'service_order': detail.pk,
                            'quantity': detail.quantity,
                            'service': service,
                            'use': use,
                            'price': detail.price,
                            'amount': detail.amount})
        context['service_conformity_detail_formset'] = \
            ServiceConformityDetailFormSet(initial=initial)
        return context



class PurchaseOrderPdfReport(View):

    def get(self, request, *args, **kwargs):
        order = PurchaseOrder.objects.get(pk=kwargs['pk'])
        response = HttpResponse(content_type='application/pdf')
        response.write(PurchaseOrderPdf().render(order))
        return response


class PurchaseOrderXlsReport(TemplateView):

    def get(self, request, *args, **kwargs):
        order = get_object_or_404(PurchaseOrder, pk=kwargs['pk'])
        wb = purchase_order_xls_report(order)
        response = HttpResponse(content_type="application/ms-excel")
        file_name = "ORDEN_DE_COMPRA_N°.xlsx"
        response["Content-Disposition"] = "attachment; filename={0}".format(file_name)
        wb.save(response)
        return response


class ServiceOrderPdfReport(View):

    def get(self, request, *args, **kwargs):
        order = ServiceOrder.objects.get(code=kwargs['code'])
        response = HttpResponse(content_type='application/pdf')
        response.write(ServiceOrderPdf().render(order))
        return response


class ServiceConformityMemoPdfReport(View):

    def get(self, request, *args, **kwargs):
        conformity = ServiceConformity.objects.get(code=kwargs['code'])
        response = HttpResponse(content_type='application/pdf')
        response.write(ServiceConformityMemoPdf().render(conformity))
        return response


class QuotationRequestPdfReport(View):

    def get(self, request, *args, **kwargs):
        quotation = Quotation.objects.get(code=kwargs['code'])
        response = HttpResponse(content_type='application/pdf')
        response.write(QuotationRequestPdf().render(quotation))
        return response


class SupplierExcelReport(TemplateView):

    def get(self, request, *args, **kwargs):
        suppliers = Supplier.objects.all().order_by('tax_id')
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE PROVEEDORES'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'RUC'
        ws['C3'] = 'RAZON_SOCIAL'
        ws['D3'] = 'DIRECCION'
        ws['E3'] = 'TELEFONO'
        ws['F3'] = 'CORREO'
        ws['G3'] = 'ESTADO_SUNAT'
        ws['H3'] = 'CONDICIÓN'
        ws['I3'] = 'REPRESENTANTE'
        ws['J3'] = 'CIIU'
        ws['K3'] = 'FECHA_ALTA'
        cont = 4
        for supplier in suppliers:
            ws.cell(row=cont, column=2).value = supplier.tax_id
            ws.cell(row=cont, column=3).value = supplier.business_name
            ws.cell(row=cont, column=4).value = supplier.address
            ws.cell(row=cont, column=5).value = supplier.phone
            ws.cell(row=cont, column=6).value = supplier.email
            ws.cell(row=cont, column=7).value = supplier.sunat_status
            ws.cell(row=cont, column=8).value = supplier.sunat_condition
            representative = supplier.representatives.first()
            ws.cell(row=cont, column=9).value = representative.name if representative else '-'
            ws.cell(row=cont, column=10).value = supplier.ciiu
            ws.cell(row=cont, column=11).value = supplier.registration_date
            ws.cell(row=cont, column=11).number_format = 'dd/mm/yyyy'
            cont = cont + 1
        file_name = "SupplierList.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class ServiceOrderExcelReportByDate(FormView):
    form_class = OrderDateReportForm
    template_name = "purchases/order_report.html"

    def form_valid(self, form):
        data = form.cleaned_data
        search_type = data['search_type']
        wb = Workbook()
        ws = wb.active
        if search_type == 'F':
            p_start_date = data['start_date']
            previous_end_date = data['end_date']
            year = int(p_start_date[6:])
            month = int(p_start_date[3:5])
            dia = int(p_start_date[0:2])
            start_date = timezone.make_aware(datetime.datetime(year, month, dia, 23, 59, 59))
            year = int(previous_end_date[6:])
            month = int(previous_end_date[3:5])
            dia = int(previous_end_date[0:2])
            end_date = timezone.make_aware(datetime.datetime(year, month, dia, 23, 59, 59))
            ws['B2'] = 'REPORTE DE ORDENES DE SERVICIOS POR FECHA'
            ws.merge_cells('B2:H2')
            ws['B3'] = 'DESDE'
            ws['C3'] = p_start_date
            ws['C3'].number_format = 'dd/mm/yyyy'
            ws['D3'] = 'HASTA'
            ws['E3'] = previous_end_date
            ws['F3'].number_format = 'dd/mm/yyyy'
            service_orders = ServiceOrder.objects.filter(date__range=[start_date, end_date])
        elif search_type == 'M':
            month = data['month'].strip()
            year = data['year'].strip()
            ws['B2'] = 'REPORTE DE ORDENES DE SERVICIOS POR MES'
            ws.merge_cells('B2:H2')
            ws['B3'] = 'MES'
            ws['C3'] = month
            ws['D3'] = 'AÑO'
            ws['E3'] = year
            service_orders = ServiceOrder.objects.filter(date__month=month, date__year=year)
        elif search_type == 'A':
            year = data['year'].strip()
            ws['B2'] = 'REPORTE DE ORDENES DE SERVICIOS POR AÑO'
            ws.merge_cells('B2:H2')
            ws['B3'] = 'AÑO'
            ws['C3'] = year
            service_orders = ServiceOrder.objects.filter(date__year=year)
        ws['B5'] = 'CODIGO'
        ws['C5'] = 'FECHA'
        ws['D5'] = 'PROVEEDOR'
        ws['E5'] = 'IMPORTE'
        ws['F5'] = 'FORMA_PAGO'
        ws['G5'] = 'CREADO'
        ws['H5'] = 'ESTADO'
        service_orders = service_orders.select_related(
            'supplier', 'payment_method', 'quotation__supplier'
        ).prefetch_related('details')
        cont = 6
        for order in service_orders:
            ws.cell(row=cont, column=2).value = order.code
            ws.cell(row=cont, column=3).value = order.date
            ws.cell(row=cont, column=3).number_format = 'dd/mm/yyyy'
            try:
                ws.cell(row=cont, column=4).value = order.quotation.supplier.business_name
            except (ObjectDoesNotExist, AttributeError):
                ws.cell(row=cont, column=4).value = order.supplier.business_name
            ws.cell(row=cont, column=5).value = order.total
            ws.cell(row=cont, column=6).value = order.payment_method.description
            ws.cell(row=cont, column=6).number_format = 'dd/mm/yyyy hh:mm:ss'
            ws.cell(row=cont, column=7).value = timezone.localtime(order.created).replace(tzinfo=None)
            ws.cell(row=cont, column=7).number_format = 'dd/mm/yyyy hh:mm:ss'
            ws.cell(row=cont, column=8).value = order.get_status_display()
            cont = cont + 1
        file_name = "ReporteOrdenesServicio.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class PurchaseOrderExcelReportByDate(FormView):
    form_class = OrderDateReportForm
    template_name = "purchases/order_report.html"

    def form_valid(self, form):
        data = form.cleaned_data
        search_type = data['search_type']
        wb = Workbook()
        ws = wb.active
        if search_type == 'F':
            p_start_date = data['start_date']
            previous_end_date = data['end_date']
            year = int(p_start_date[6:])
            month = int(p_start_date[3:5])
            dia = int(p_start_date[0:2])
            start_date = timezone.make_aware(datetime.datetime(year, month, dia, 23, 59, 59))
            year = int(previous_end_date[6:])
            month = int(previous_end_date[3:5])
            dia = int(previous_end_date[0:2])
            end_date = timezone.make_aware(datetime.datetime(year, month, dia, 23, 59, 59))
            ws['B2'] = 'REPORTE DE ORDENES DE COMPRA POR FECHA'
            ws.merge_cells('B2:H2')
            ws['B3'] = 'DESDE'
            ws['C3'] = p_start_date
            ws['C3'].number_format = 'dd/mm/yyyy'
            ws['D3'] = 'HASTA'
            ws['E3'] = previous_end_date
            ws['F3'].number_format = 'dd/mm/yyyy'
            purchase_orders = PurchaseOrder.objects.filter(date__range=[start_date, end_date])
        elif search_type == 'M':
            month = data['month'].strip()
            year = data['year'].strip()
            ws['B2'] = 'REPORTE DE ORDENES DE COMPRA POR MES'
            ws.merge_cells('B2:H2')
            ws['B3'] = 'MES'
            ws['C3'] = month
            ws['D3'] = 'AÑO'
            ws['E3'] = year
            purchase_orders = PurchaseOrder.objects.filter(date__month=month, date__year=year)
        elif search_type == 'A':
            year = data['year'].strip()
            ws['B2'] = 'REPORTE DE ORDENES DE COMPRA POR AÑO'
            ws.merge_cells('B2:H2')
            ws['B3'] = 'AÑO'
            ws['C3'] = year
            purchase_orders = PurchaseOrder.objects.filter(date__year=year)
        ws['B5'] = 'CODIGO'
        ws['C5'] = 'FECHA'
        ws['D5'] = 'PROVEEDOR'
        ws['E5'] = 'IMPORTE'
        ws['F5'] = 'FORMA_PAGO'
        ws['G5'] = 'CREADO'
        ws['H5'] = 'ESTADO'
        purchase_orders = purchase_orders.select_related(
            'supplier', 'payment_method', 'quotation__supplier'
        ).prefetch_related('details')
        cont = 6
        for purchase_order in purchase_orders:
            ws.cell(row=cont, column=2).value = purchase_order.code
            ws.cell(row=cont, column=3).value = purchase_order.date
            ws.cell(row=cont, column=3).number_format = 'dd/mm/yyyy'
            try:
                ws.cell(row=cont, column=4).value = purchase_order.quotation.supplier.business_name
            except (ObjectDoesNotExist, AttributeError):
                ws.cell(row=cont, column=4).value = purchase_order.supplier.business_name
            ws.cell(row=cont, column=5).value = purchase_order.total
            ws.cell(row=cont, column=6).value = purchase_order.payment_method.description
            ws.cell(row=cont, column=6).number_format = 'dd/mm/yyyy hh:mm:ss'
            ws.cell(row=cont, column=7).value = timezone.localtime(purchase_order.created).replace(tzinfo=None)
            ws.cell(row=cont, column=7).number_format = 'dd/mm/yyyy hh:mm:ss'
            ws.cell(row=cont, column=8).value = purchase_order.get_status_display()
            cont = cont + 1
        file_name = "ReporteOrdenesCompra.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class QuotationTransfer(TemplateView):
    template_name = 'purchases/quotation_transfer.html'

    def get_context_data(self, **kwargs):
        context = super(QuotationTransfer, self).get_context_data(**kwargs)
        context['quotations'] = Quotation.objects.filter(status=Quotation.STATUS.PEND)
        return context


class PurchaseOrderTransfer(TemplateView):
    template_name = 'purchases/purchase_order_transfer.html'

    def get_context_data(self, **kwargs):
        context = super(PurchaseOrderTransfer, self).get_context_data(**kwargs)
        context['orders'] = PurchaseOrder.objects.filter(
            Q(status=PurchaseOrder.STATUS.PEND) | Q(status=PurchaseOrder.STATUS.ING_PARC))
        return context


class ServiceOrderTransfer(TemplateView):
    template_name = 'purchases/service_order_transfer.html'

    def get_context_data(self, **kwargs):
        context = super(ServiceOrderTransfer, self).get_context_data(**kwargs)
        context['orders'] = ServiceOrder.objects.filter(status=ServiceOrder.STATUS.PEND)
        return context
