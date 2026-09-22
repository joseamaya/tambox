# -*- coding: utf-8 -*-
from django.shortcuts import render, get_object_or_404
from django.views.generic.base import View, TemplateView
from django.views.generic.list import ListView
from django.views.generic.edit import FormView, UpdateView, CreateView
from django.urls import reverse_lazy, reverse
from django.http.response import HttpResponseRedirect
import json
from django.http import HttpResponse
import simplejson
from django.views.generic.detail import DetailView
from administration.models import Office, ApprovalLevel
import locale
from security.permissions import requires
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.db import transaction, IntegrityError
from django.core.exceptions import ObjectDoesNotExist
from django.db.models import Q
from django.contrib import messages
from requirements.models import RequirementApproval, Requirement,\
    RequirementDetail
from requirements.forms import RequirementApprovalForm, RequirementForm, RequirementDetailFormSet
from purchases.forms import QuotationDetailFormSet
from purchases.models import Quotation
from products.models import Product
from requirements.mail import requirement_creation_mail
from openpyxl import Workbook
from requirements.reports import RequirementReport
from datetime import date
from tambox.config import configuration, administration_office,\
    logistics, budget

locale.setlocale(locale.LC_ALL, "")


from tambox.views import AjaxOnlyMixin, HtmxListMixin
class Dashboard(View):
    def get(self, request, *args, **kwargs):
        notification_list = []
        context = {'notifications': notification_list}
        return render(request, 'requirements/requirements_dashboard.html', context)


class RequirementApprove(UpdateView):
    model = RequirementApproval
    template_name = 'requirements/requirement_approve.html'
    form_class = RequirementApprovalForm
    success_url = reverse_lazy('requirements:requirement_approval_list')

    @method_decorator(requires('requirements.change_requirementapproval'))
    def dispatch(self, *args, **kwargs):
        requirement_approval = get_object_or_404(self.model, pk=kwargs['pk'])
        user = self.request.user
        if requirement_approval.check_approval_access(user):
            return super(RequirementApprove, self).dispatch(*args, **kwargs)
        else:
            return HttpResponseRedirect(reverse('security:permission_denied'))

    def get_form_kwargs(self):
        kwargs = super(RequirementApprove, self).get_form_kwargs()
        kwargs['request'] = self.request
        return kwargs

    def form_valid(self, form):
        form.save()
        return HttpResponseRedirect(self.get_success_url())

    def form_invalid(self, form):
        return self.render_to_response(self.get_context_data(form=form))


class RequirementDetailCreate(AjaxOnlyMixin, FormView):
    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            detail_list = []
            det = {}
            det['code'] = ''
            det['product'] = ''
            det['unit'] = ''
            det['quantity'] = '0'
            det['use'] = ''
            detail_list.append(det)
            formset = RequirementDetailFormSet(initial=detail_list)
            json_list = []
            for form in formset:
                detail_json = {}
                detail_json['code'] = str(form['code'])
                detail_json['product'] = str(form['product'])
                detail_json['unit'] = str(form['unit'])
                detail_json['quantity'] = str(form['quantity'])
                detail_json['use'] = str(form['use'])
                json_list.append(detail_json)
            data = json.dumps(json_list)
            return HttpResponse(data, 'application/json')


class RequirementCreate(CreateView):
    template_name = 'requirements/requirement_form.html'
    form_class = RequirementForm
    model = Requirement
    context_object_name = 'requirement'

    def get_form_kwargs(self):
        kwargs = super(RequirementCreate, self).get_form_kwargs()
        kwargs['request'] = self.request
        return kwargs

    def get_initial(self):
        initial = super(RequirementCreate, self).get_initial()
        initial['year'] = date.today().year
        initial['date'] = date.today().strftime('%d/%m/%Y')
        initial['month'] = date.today().month
        return initial

    def get(self, request, *args, **kwargs):
        self.object = None
        offices = Office.objects.all()
        if not offices:
            return HttpResponseRedirect(reverse('administration:office_create'))
        try:
            worker = self.request.user.worker
        except ObjectDoesNotExist:
            return HttpResponseRedirect(reverse('administration:worker_create'))
        if worker.signature == '':
            return HttpResponseRedirect(reverse('administration:worker_update', args=[worker.pk]))
        position = worker.position
        if position is None:
            return HttpResponseRedirect(reverse('administration:position_create'))
        boss_position = position.superior_position
        if boss_position is None:
            return HttpResponseRedirect(reverse('administration:position_create'))
        approval_levels = ApprovalLevel.objects.all()
        if not approval_levels:
            return HttpResponseRedirect(reverse('administration:approval_level_create'))
        if configuration() is not None:
            form_class = self.get_form_class()
            form = self.get_form(form_class)
            requirement_detail_formset = RequirementDetailFormSet()
            return self.render_to_response(self.get_context_data(form=form,
                                                                 requirement_detail_formset=requirement_detail_formset))
        else:
            return HttpResponseRedirect(reverse('accounting:configuration'))

    def post(self, request, *args, **kwargs):
        self.object = None
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        requirement_detail_formset = RequirementDetailFormSet(request.POST)
        if form.is_valid() and requirement_detail_formset.is_valid():
            return self.form_valid(form, requirement_detail_formset)
        else:
            return self.form_invalid(form, requirement_detail_formset)

    def form_valid(self, form, requirement_detail_formset):
        try:
            with transaction.atomic():
                self.object = form.save()
                details = []
                cont = 1
                for requirement_detail_form in requirement_detail_formset:
                    code = requirement_detail_form.cleaned_data.get('code')
                    quantity = requirement_detail_form.cleaned_data.get('quantity')
                    use = requirement_detail_form.cleaned_data.get('use')
                    if quantity:
                        if code:
                            detail = RequirementDetail(requirement=self.object,
                                                           line_number=cont,
                                                           product=Product.objects.get(code=code),
                                                           quantity=quantity,
                                                           use=use)
                        else:
                            detail = RequirementDetail(requirement=self.object,
                                                           line_number=cont,
                                                           otro=requirement_detail_form.cleaned_data.get('product'),
                                                           quantity=quantity,
                                                           use=use)
                        details.append(detail)
                        cont = cont + 1
                RequirementDetail.objects.bulk_create(details)
                boss_position = self.object.requester.position.superior_position  # Position.objects.get(office=self.object.office, is_leadership=True, is_active=True)
                boss = boss_position.worker
                destinatario = boss.user.email
                if boss.pk != self.object.requester.pk:
                    requirement_creation_mail(destinatario, self.object)
                return HttpResponseRedirect(reverse('requirements:requirement_detail', args=[self.object.code]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando el requerimiento.')

    def form_invalid(self, form, requirement_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             requirement_detail_formset=requirement_detail_formset))


class RequirementDetailView(DetailView):
    model = Requirement
    context_object_name = 'requirement'
    slug_field = 'code'
    slug_url_kwarg = 'code'
    template_name = 'requirements/requirement_detail.html'

    @method_decorator(
        requires('requirements.ver_detalle_requerimiento'))
    def dispatch(self, *args, **kwargs):
        requirement = self.get_object()
        if requirement.check_access(self.request.user, administration_office(), logistics(), budget()):
            return super(RequirementDetailView, self).dispatch(*args, **kwargs)
        else:
            return HttpResponseRedirect(
                reverse('requirements:requirement_detail', args=[requirement.next()]))


class RequirementDelete(TemplateView):
    http_method_names = ['post']

    @method_decorator(requires('requirements.delete_requirement'))
    def dispatch(self, *args, **kwargs):
        return super(RequirementDelete, self).dispatch(*args, **kwargs)

    def post(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            code = request.POST['code']
            requirement = Requirement.objects.get(code=code)
            requirement_json = {}
            requirement_json['code'] = code
            quotations = requirement.quotations.all()
            if len(quotations) > 0:
                requirement_json['quotations'] = 'SI'
            else:
                requirement_json['quotations'] = 'NO'
                with transaction.atomic():
                    requirement.delete_requirement()
                    RequirementDetail.objects.filter(requirement=requirement).delete()
            data = simplejson.dumps(requirement_json)
            return HttpResponse(data, 'application/json')


class RequirementApprovalList(HtmxListMixin, ListView):
    model = RequirementApproval
    template_name = 'requirements/requirement_approval_list.html'
    fragment_template_name = 'requirements/includes/requirement_approval_rows.html'
    context_object_name = 'requirement_approvals'
    paginate_by = 10
    search_fields = ('requirement__code',)

    @method_decorator(
        requires('requirements.ver_tabla_requerimientos'))
    def dispatch(self, *args, **kwargs):
        return super(RequirementApprovalList, self).dispatch(*args, **kwargs)

    def get(self, request, *args, **kwargs):
        try:
            worker = self.request.user.worker
        except ObjectDoesNotExist:
            return HttpResponseRedirect(reverse('administration:worker_create'))
        if worker.signature == '':
            return HttpResponseRedirect(reverse('administration:worker_update', args=[worker.pk]))
        position = worker.position
        if position is None:
            return HttpResponseRedirect(reverse('administration:position_create'))
        if not position.is_leadership:
            return HttpResponseRedirect(reverse('security:permission_denied'))
        return super(RequirementApprovalList, self).get(request, *args, **kwargs)

    def get_base_queryset(self):
        return RequirementApproval.get_pending_approvals(self.request.user)


class QuotationListByRequirement(ListView):
    model = Quotation
    template_name = 'purchases/quotation_list.html'
    context_object_name = 'quotations'

    @method_decorator(requires('purchases.ver_tabla_cotizaciones'))
    def dispatch(self, *args, **kwargs):
        return super(QuotationListByRequirement, self).dispatch(*args, **kwargs)

    def get_queryset(self):
        requirement = Requirement.objects.get(pk=self.kwargs['requirement'])
        queryset = requirement.quotations.all()
        return queryset


class RequirementList(HtmxListMixin, ListView):
    model = Requirement
    template_name = 'requirements/requirement_list.html'
    fragment_template_name = 'requirements/includes/requirement_rows.html'
    context_object_name = 'requirements'
    paginate_by = 10
    search_fields = ('code', 'office__name')

    def get_base_queryset(self):
        user = self.request.user
        return Requirement.get_visible_requirements(user)

    @method_decorator(
        requires('requirements.ver_tabla_requerimientos'))
    def dispatch(self, *args, **kwargs):
        return super(RequirementList, self).dispatch(*args, **kwargs)


class RequirementUpdate(UpdateView):
    template_name = 'requirements/requirement_form.html'
    model = Requirement
    context_object_name = 'requirement'
    form_class = RequirementForm

    @method_decorator(
        requires('requirements.change_requirement'))
    def dispatch(self, *args, **kwargs):
        requirement = self.get_object()
        if requirement.approval.is_active or self.request.user.is_superuser:
            return super(RequirementUpdate, self).dispatch(*args, **kwargs)
        else:
            return HttpResponseRedirect(reverse('security:permission_denied'))

    def get_form_kwargs(self):
        kwargs = super(RequirementUpdate, self).get_form_kwargs()
        kwargs['request'] = self.request
        return kwargs

    def get_initial(self):
        initial = super(RequirementUpdate, self).get_initial()
        initial['date'] = self.object.date.strftime('%d/%m/%Y')
        return initial

    def get_context_data(self, **kwargs):
        context = super(RequirementUpdate, self).get_context_data(**kwargs)
        return context

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        details = RequirementDetail.objects.filter(requirement=self.object).order_by('line_number')
        details_data = []
        for detail in details:
            if detail.product is None:
                d = {'code': '',
                     'product': detail.otro or '',
                     'quantity': detail.quantity,
                     'unit': '',
                     'use': detail.use}
            else:
                d = {'code': detail.product.code,
                     'product': detail.product.description,
                     'quantity': detail.quantity,
                     'unit': detail.product.unit_of_measure.code,
                     'use': detail.use}
            details_data.append(d)
        requirement_detail_formset = RequirementDetailFormSet(initial=details_data)
        return self.render_to_response(self.get_context_data(form=form,
                                                             requirement_detail_formset=requirement_detail_formset))

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        form_class = self.get_form_class()
        form = self.get_form(form_class)
        requirement_detail_formset = RequirementDetailFormSet(request.POST)
        if form.is_valid() and requirement_detail_formset.is_valid():
            return self.form_valid(form, requirement_detail_formset)
        else:
            return self.form_invalid(form, requirement_detail_formset)

    def form_valid(self, form, requirement_detail_formset):
        try:
            with transaction.atomic():
                RequirementDetail.objects.filter(requirement=self.object).delete()
                form.save()
                details = []
                cont = 1
                for requirement_detail_form in requirement_detail_formset:
                    code = requirement_detail_form.cleaned_data.get('code')
                    quantity = requirement_detail_form.cleaned_data.get('quantity')
                    use = requirement_detail_form.cleaned_data.get('use')
                    if quantity:
                        if code:
                            detail = RequirementDetail(requirement=self.object,
                                                           line_number=cont,
                                                           product=Product.objects.get(code=code),
                                                           quantity=quantity,
                                                           use=use)
                        else:
                            detail = RequirementDetail(requirement=self.object,
                                                           line_number=cont,
                                                           otro=requirement_detail_form.cleaned_data.get('product'),
                                                           quantity=quantity,
                                                           use=use)
                        details.append(detail)
                        cont = cont + 1
                RequirementDetail.objects.bulk_create(details)
                return HttpResponseRedirect(reverse('requirements:requirement_detail', args=[self.object.code]))
        except IntegrityError:
            messages.error(self.request, 'Error guardando el requerimiento.')

    def form_invalid(self, form, requirement_detail_formset):
        return self.render_to_response(self.get_context_data(form=form,
                                                             requirement_detail_formset=requirement_detail_formset))


class RequirementDetailFetch(AjaxOnlyMixin, TemplateView):

    required_params = ('requirement', 'search_type')
    def get(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            requirement = request.GET['requirement']
            search_type = request.GET['search_type']
            if search_type == 'TODOS':
                details = RequirementDetail.objects.filter(
                    Q(status=RequirementDetail.STATUS.PEND) | Q(status=RequirementDetail.STATUS.COTIZ),
                    requirement__code=requirement).order_by('line_number')
            elif search_type == 'PRODUCTOS':
                details = RequirementDetail.objects.filter(Q(status=RequirementDetail.STATUS.PEND) |
                                                               Q(status=RequirementDetail.STATUS.COTIZ),
                                                               requirement__code=requirement,
                                                               product__isnull=False).order_by('line_number')
            detail_list = []
            for detail in details:
                det = {}
                det['requirement'] = detail.id
                try:
                    det['code'] = detail.product.code
                    det['name'] = detail.product.description
                    det['unit'] = detail.product.unit_of_measure.code
                    # det['use'] = detail.use
                    det['quantity'] = str(detail.quantity - detail.served_quantity)
                    # det['price'] = str(detail.product.price)
                    # det['amount'] = str(detail.product.price*(detail.quantity-detail.served_quantity))
                    detail_list.append(det)
                except AttributeError:
                    pass
            formset = QuotationDetailFormSet(initial=detail_list)
            json_list = []
            for form in formset:
                detail_json = {}
                detail_json['requirement'] = str(form['requirement'])
                detail_json['code'] = str(form['code'])
                detail_json['name'] = str(form['name'])
                detail_json['unit'] = str(form['unit'])
                detail_json['quantity'] = str(form['quantity'])
                json_list.append(detail_json)
            data = json.dumps(json_list)
            return HttpResponse(data, 'application/json')


class RequirementTransfer(TemplateView):
    template_name = 'requirements/requirement_transfer.html'

    def get_context_data(self, **kwargs):
        context = super(RequirementTransfer, self).get_context_data(**kwargs)
        # requerimientos = Requirement.objects.all()
        requirements = Requirement.get_requirements_ready_for_transfer()
        context['requirements'] = requirements
        return context


class RequirementExcelReport(TemplateView):
    def get(self, request, *args, **kwargs):
        requirements = Requirement.objects.active_requirements_by_user(request.user,
                                                                                  Requirement.STATUS.CANC)
        wb = Workbook()
        ws = wb.active
        ws['B1'] = 'REPORTE DE REQUERIMIENTOS'
        ws.merge_cells('B1:J1')
        ws['B3'] = 'CODIGO'
        ws['C3'] = 'OFICINA'
        ws['D3'] = 'ESTADO APROBACION'
        ws['E3'] = 'ESTADO'
        ws['F3'] = 'FECHA'
        cont = 4
        for requirement in requirements:
            ws.cell(row=cont, column=2).value = requirement.code
            ws.cell(row=cont, column=3).value = requirement.office.name
            ws.cell(row=cont, column=4).value = requirement.approval.level.description
            ws.cell(row=cont, column=5).value = requirement.get_status_display()
            ws.cell(row=cont, column=6).value = timezone.localtime(requirement.created).replace(tzinfo=None)
            ws.cell(row=cont, column=6).number_format = 'dd/mm/yyyy hh:mm:ss'
            cont = cont + 1
        file_name = "RequirementList.xlsx"
        response = HttpResponse(content_type="application/ms-excel")
        contenido = "attachment; filename={0}".format(file_name)
        response["Content-Disposition"] = contenido
        wb.save(response)
        return response


class RequirementPdfReport(View):
    def get(self, request, *args, **kwargs):
        code = kwargs['code']
        requirement = Requirement.objects.get(code=code)
        response = HttpResponse(content_type='application/pdf')
        report = RequirementReport('A4', requirement)
        pdf = report.render()
        response.write(pdf)
        return response
